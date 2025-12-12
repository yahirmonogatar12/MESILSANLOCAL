# -*- coding: utf-8 -*-
"""
Rutas de Producción y Plan de Producción
APIs para gestión de planes de producción, work orders, embarque y control de línea
"""

import io
import traceback
from datetime import datetime, timedelta
from flask import Blueprint, request, jsonify, session, render_template, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from .utils import login_requerido, obtener_fecha_hora_mexico
from ..database.db_mysql import execute_query, get_mysql_connection
from ..database.db import get_db_connection

produccion_bp = Blueprint('produccion', __name__)


# ============== HELPERS ==============

def _fp_safe_date(val):
    """Convertir valor a fecha de forma segura"""
    if not val:
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, str):
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%Y/%m/%d'):
            try:
                return datetime.strptime(val, fmt).date()
            except:
                continue
    return None


def _fp_generate_lot_no(working_date):
    """Generar número de lote único para un plan"""
    prefix = 'ASSYLINE'
    fecha_str = working_date.strftime('%y%m%d')
    
    # Obtener el siguiente secuencial del día
    seq_query = """
        SELECT MAX(CAST(SUBSTRING_INDEX(lot_no, '-', -1) AS UNSIGNED)) as max_seq 
        FROM plan_main 
        WHERE lot_no LIKE %s
    """
    pattern = f"{prefix}-{fecha_str}-%"
    result = execute_query(seq_query, (pattern,), fetch='one')
    max_seq = result.get('max_seq') if result else None
    next_seq = (max_seq or 0) + 1
    
    return f"{prefix}-{fecha_str}-{next_seq:03d}"


# ============== API PLAN MAIN ==============

@produccion_bp.route('/api/plan', methods=['GET'])
@login_requerido
def api_plan_list():
    """Listar planes de producción con filtros de fecha"""
    try:
        start = request.args.get('start')
        end = request.args.get('end')
        where = []
        params = []
        
        if start:
            if not end:
                where.append('DATE(working_date) = %s')
                params.append(start)
            else:
                where.append('DATE(working_date) >= %s')
                params.append(start)
        if end:
            where.append('DATE(working_date) <= %s')
            params.append(end)
            
        sql = (
            "SELECT id, lot_no, wo_code, po_code, working_date, line, routing, model_code, part_no, project, process, "
            "COALESCE(ct,0) AS ct, COALESCE(uph,0) AS uph, COALESCE(plan_count,0) AS plan_count, "
            "COALESCE(produced_count,0) AS input, 0 AS output, COALESCE(entregadas_main,0) AS entregadas_main, "
            "COALESCE(produced_count,0) AS produced, status, group_no, sequence FROM plan_main"
        )
        
        if where:
            sql += ' WHERE ' + ' AND '.join(where)
        sql += ' ORDER BY COALESCE(group_no,999), COALESCE(sequence,999), working_date, created_at'
        
        rows = execute_query(sql, tuple(params) if params else None, fetch='all')
        
        data = []
        for r in rows:
            data.append({
                'lot_no': r.get('lot_no') if isinstance(r, dict) else r[1],
                'wo_code': r.get('wo_code') if isinstance(r, dict) else r[2],
                'po_code': r.get('po_code') if isinstance(r, dict) else r[3],
                'working_date': str((r.get('working_date') if isinstance(r, dict) else r[4]) or '')[:10],
                'line': r.get('line') if isinstance(r, dict) else r[5],
                'routing': r.get('routing') if isinstance(r, dict) else r[6],
                'model_code': r.get('model_code') if isinstance(r, dict) else r[7],
                'part_no': r.get('part_no') if isinstance(r, dict) else r[8],
                'project': r.get('project') if isinstance(r, dict) else r[9],
                'process': r.get('process') if isinstance(r, dict) else r[10],
                'ct': r.get('ct') if isinstance(r, dict) else r[11],
                'uph': r.get('uph') if isinstance(r, dict) else r[12],
                'plan_count': r.get('plan_count') if isinstance(r, dict) else r[13],
                'input': r.get('input') if isinstance(r, dict) else r[14],
                'output': r.get('output') if isinstance(r, dict) else r[15],
                'entregadas_main': r.get('entregadas_main') if isinstance(r, dict) else r[16],
                'produced': r.get('produced') if isinstance(r, dict) else r[17],
                'status': r.get('status') if isinstance(r, dict) else r[18],
                'group_no': r.get('group_no') if isinstance(r, dict) else r[19],
                'sequence': r.get('sequence') if isinstance(r, dict) else r[20],
            })
        return jsonify(data)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@produccion_bp.route('/api/plan', methods=['POST'])
@login_requerido
def api_plan_create():
    """Crear un nuevo plan de producción"""
    try:
        data = request.get_json() or {}
        working_date = data.get('working_date')
        part_no = data.get('part_no')
        line = data.get('line')
        turno = (data.get('turno') or 'DIA').strip().upper()
        plan_count = int(data.get('plan_count') or 0)
        
        wo_code = data.get('wo_code') or 'SIN-WO'
        po_code = data.get('po_code') or 'SIN-PO'
        
        if not (working_date and part_no and line):
            return jsonify({'error': 'Parámetros requeridos'}), 400
            
        fecha = _fp_safe_date(working_date) or datetime.utcnow().date()
        routing = {'DIA': 1, 'TIEMPO EXTRA': 2, 'NOCHE': 3}.get(turno, 1)
        lot_no = _fp_generate_lot_no(datetime.combine(fecha, datetime.min.time()))
        
        # Buscar información adicional en raw
        raw_data_query = """
            SELECT part_no, model, project, c_t as ct, uph
            FROM raw
            WHERE part_no = %s OR part_no LIKE %s OR model = %s OR model LIKE %s
            ORDER BY id DESC
            LIMIT 1
        """
        raw_data = execute_query(raw_data_query, (part_no, f"%{part_no}%", part_no, f"%{part_no}%"), fetch='one')
        
        if raw_data:
            model_code = raw_data.get('model') or part_no
            project = raw_data.get('project') or ''
            try:
                ct = float(raw_data.get('ct') or 0)
            except:
                ct = 0.0
            try:
                uph_raw = raw_data.get('uph')
                uph = int(float(str(uph_raw).strip())) if uph_raw and str(uph_raw).strip().replace('.', '').isdigit() else 0
            except:
                uph = 0
        else:
            model_code = part_no
            project = ''
            ct = 0.0
            uph = 0
        
        group_no = data.get('group_no')
        sequence = None
        
        if group_no is not None:
            seq_query = "SELECT MAX(sequence) as max_seq FROM plan_main WHERE group_no = %s"
            seq_result = execute_query(seq_query, (int(group_no),), fetch='one')
            max_seq = seq_result.get('max_seq') if seq_result else None
            sequence = (max_seq + 1) if max_seq is not None else 1
        
        if group_no is not None and sequence is not None:
            sql = (
                "INSERT INTO plan_main (lot_no, wo_code, po_code, working_date, line, model_code, part_no, project, process, plan_count, ct, uph, routing, status, group_no, sequence, created_at) "
                "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'PLAN',%s,%s,NOW())"
            )
            params = (lot_no, wo_code, po_code, fecha, line, model_code, part_no, project, 'MAIN', plan_count, ct, uph, routing, int(group_no), sequence)
        else:
            sql = (
                "INSERT INTO plan_main (lot_no, wo_code, po_code, working_date, line, model_code, part_no, project, process, plan_count, ct, uph, routing, status, created_at) "
                "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'PLAN',NOW())"
            )
            params = (lot_no, wo_code, po_code, fecha, line, model_code, part_no, project, 'MAIN', plan_count, ct, uph, routing)
        
        execute_query(sql, params)
        return jsonify({'success': True, 'lot_no': lot_no, 'model_code': model_code, 'ct': ct, 'uph': uph, 'project': project})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@produccion_bp.route('/api/plan/update', methods=['POST'])
@login_requerido
def api_plan_update():
    """Actualizar un plan existente"""
    try:
        data = request.get_json() or {}
        lot_no = data.get('lot_no')
        if not lot_no:
            return jsonify({'error': 'lot_no requerido'}), 400
            
        fields = []
        vals = []
        
        if 'plan_count' in data:
            fields.append('plan_count = %s')
            vals.append(int(data.get('plan_count') or 0))
        if 'status' in data:
            fields.append('status = %s')
            vals.append(str(data.get('status')))
        if 'line' in data:
            fields.append('line = %s')
            vals.append(str(data.get('line')))
        if 'wo_code' in data:
            fields.append('wo_code = %s')
            vals.append(str(data.get('wo_code')))
        if 'po_code' in data:
            fields.append('po_code = %s')
            vals.append(str(data.get('po_code')))
        if 'turno' in data:
            routing = {'DIA': 1, 'TIEMPO EXTRA': 2, 'NOCHE': 3}.get(str(data.get('turno')).strip().upper(), 1)
            fields.append('routing = %s')
            vals.append(routing)
        if 'uph' in data:
            fields.append('uph = %s')
            vals.append(str(data.get('uph')))
        if 'ct' in data:
            fields.append('ct = %s')
            vals.append(str(data.get('ct')))
        if 'project' in data:
            fields.append('project = %s')
            vals.append(str(data.get('project')))
        if 'model_code' in data:
            fields.append('model_code = %s')
            vals.append(str(data.get('model_code')))
            
        if not fields:
            return jsonify({'error': 'Sin cambios'}), 400
            
        fields.append('updated_at = NOW()')
        sql = f"UPDATE plan_main SET {', '.join(fields)} WHERE lot_no = %s"
        vals.append(lot_no)
        execute_query(sql, tuple(vals))
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@produccion_bp.route('/api/plan/status', methods=['POST'])
@login_requerido
def api_plan_status():
    """Actualizar el status de un plan con validaciones"""
    try:
        data = request.get_json() or {}
        lot_no = data.get('lot_no', '').strip()
        new_status = data.get('status', '').strip().upper()
        
        if not lot_no:
            return jsonify({'error': 'lot_no requerido', 'error_code': 'MISSING_LOT_NO'}), 400
        
        if not new_status:
            return jsonify({'error': 'status requerido', 'error_code': 'MISSING_STATUS'}), 400
        
        valid_statuses = ['PENDIENTE', 'EN PROGRESO', 'PAUSADO', 'TERMINADO', 'CANCELADO']
        if new_status not in valid_statuses:
            return jsonify({'error': f'Status inválido: {new_status}', 'error_code': 'INVALID_STATUS'}), 400
        
        check_sql = "SELECT line, status, plan_count, produced_count, started_at, pause_started_at, paused_at FROM plan_main WHERE lot_no = %s"
        plan_result = execute_query(check_sql, (lot_no,), fetch='one')
        
        if not plan_result:
            return jsonify({'error': 'Plan no encontrado', 'error_code': 'NOT_FOUND'}), 404
        
        current_line = plan_result.get('line')
        current_status = (plan_result.get('status') or '').strip().upper()
        plan_count = int(plan_result.get('plan_count') or 0)
        produced_count = int(plan_result.get('produced_count') or 0)
        started_at = plan_result.get('started_at')
        pause_started_at = plan_result.get('pause_started_at')
        paused_at = int(plan_result.get('paused_at') or 0)
        
        # Validación de conflicto de línea
        if new_status == 'EN PROGRESO' and current_status != 'EN PROGRESO':
            conflict_sql = "SELECT lot_no FROM plan_main WHERE line = %s AND status = 'EN PROGRESO' AND lot_no != %s LIMIT 1"
            conflict_result = execute_query(conflict_sql, (current_line, lot_no), fetch='one')
            
            if conflict_result:
                return jsonify({
                    'error': 'Ya existe un plan EN PROGRESO en esta línea',
                    'error_code': 'LINE_CONFLICT',
                    'line': current_line,
                    'lot_no_en_progreso': conflict_result.get('lot_no')
                }), 409
        
        update_fields = ['status = %s', 'updated_at = NOW()']
        update_values = [new_status]
        
        if new_status == 'EN PROGRESO':
            if current_status == 'PAUSADO' and pause_started_at:
                update_fields.append('paused_at = paused_at + TIMESTAMPDIFF(SECOND, pause_started_at, NOW())')
            elif current_status != 'EN PROGRESO' and not started_at:
                update_fields.append('started_at = NOW()')
        
        if new_status == 'PAUSADO' and current_status == 'EN PROGRESO':
            if 'pause_reason' in data:
                update_fields.append('pause_reason = %s')
                update_values.append(str(data.get('pause_reason', '')))
            update_fields.append('pause_started_at = NOW()')
        
        if new_status == 'TERMINADO':
            if current_status == 'PAUSADO' and pause_started_at:
                update_fields.append('paused_at = paused_at + TIMESTAMPDIFF(SECOND, pause_started_at, NOW())')
            update_fields.append('ended_at = NOW()')
            if produced_count < plan_count and 'end_reason' in data:
                update_fields.append('end_reason = %s')
                update_values.append(str(data.get('end_reason', '')))
        
        update_sql = f"UPDATE plan_main SET {', '.join(update_fields)} WHERE lot_no = %s"
        update_values.append(lot_no)
        
        rows_affected = execute_query(update_sql, tuple(update_values))
        
        if isinstance(rows_affected, int) and rows_affected == 0:
            return jsonify({'error': 'No se actualizó ninguna fila', 'error_code': 'NO_ROWS_UPDATED'}), 400
        
        return jsonify({
            'success': True,
            'lot_no': lot_no,
            'new_status': new_status,
            'line': current_line
        })
        
    except Exception as e:
        print(f"Error en api_plan_status: {e}")
        traceback.print_exc()
        return jsonify({'error': str(e), 'error_code': 'UNHANDLED_EXCEPTION'}), 500


@produccion_bp.route('/api/plan/save-sequences', methods=['POST'])
@login_requerido
def api_plan_save_sequences():
    """Guardar secuencias de planes"""
    try:
        payload = request.get_json() or {}
        sequences = payload.get('sequences', [])
        updated = 0
        
        for item in sequences:
            lot_no = item.get('lot_no')
            group_no = item.get('group_no')
            sequence = item.get('sequence')
            
            if not (lot_no and group_no is not None and sequence is not None):
                continue
                
            vals = []
            sets = []
            sets.append('group_no = %s'); vals.append(int(group_no))
            sets.append('sequence = %s'); vals.append(int(sequence))
            
            if item.get('plan_start_date') and item.get('plan_start_date') != '--':
                sets.append('plan_start_date = %s'); vals.append(item['plan_start_date'])
            if item.get('planned_start') and item.get('planned_start') != '--':
                sets.append('planned_start = %s'); vals.append(item['planned_start'])
            if item.get('planned_end') and item.get('planned_end') != '--':
                sets.append('planned_end = %s'); vals.append(item['planned_end'])
            if 'effective_minutes' in item:
                sets.append('effective_minutes = %s'); vals.append(int(item.get('effective_minutes') or 0))
            if 'breaks_minutes' in item:
                sets.append('breaks_minutes = %s'); vals.append(int(item.get('breaks_minutes') or 0))
            
            sets.append('updated_at = NOW()')
            vals.append(lot_no)
            sql = f"UPDATE plan_main SET {', '.join(sets)} WHERE lot_no = %s"
            execute_query(sql, tuple(vals))
            updated += 1
            
        return jsonify({'success': True, 'updated_count': updated, 'message': f'{updated} secuencias guardadas correctamente'})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@produccion_bp.route('/api/plan/pending', methods=['GET'])
@login_requerido
def api_plan_pending():
    """Obtener planes con cantidad pendiente"""
    try:
        start = request.args.get('start')
        end = request.args.get('end')
        
        where = ["status <> 'CANCELADO'"]
        params = []
        
        if start:
            where.append('DATE(working_date) >= %s')
            params.append(start)
        if end:
            where.append('DATE(working_date) <= %s')
            params.append(end)
        
        where.append('COALESCE(plan_count, 0) > COALESCE(produced_count, 0)')
        
        sql = (
            "SELECT lot_no, working_date, part_no, line, "
            "COALESCE(plan_count,0) AS plan_count, "
            "COALESCE(produced_count,0) AS input, status "
            "FROM plan_main WHERE " + ' AND '.join(where) + " ORDER BY working_date, lot_no"
        )
        
        rows = execute_query(sql, tuple(params) if params else None, fetch='all')
        
        data = []
        for r in rows:
            data.append({
                'lot_no': r['lot_no'] if isinstance(r, dict) else r[0],
                'working_date': str((r['working_date'] if isinstance(r, dict) else r[1]) or '')[:10],
                'part_no': r['part_no'] if isinstance(r, dict) else r[2],
                'line': r['line'] if isinstance(r, dict) else r[3],
                'plan_count': r['plan_count'] if isinstance(r, dict) else r[4],
                'input': r['input'] if isinstance(r, dict) else r[5],
                'status': r['status'] if isinstance(r, dict) else r[6]
            })
        
        return jsonify(data)
        
    except Exception as e:
        print(f"❌ Error en api_plan_pending: {str(e)}")
        return jsonify({'error': str(e)}), 500


@produccion_bp.route('/api/plan/reschedule', methods=['POST'])
@login_requerido
def api_plan_reschedule():
    """Reprogramar planes pendientes creando nuevos planes"""
    try:
        data = request.get_json() or {}
        lot_nos = data.get('lot_nos', [])
        new_date = data.get('new_working_date')
        
        if not (lot_nos and new_date):
            return jsonify({'error': 'Parámetros requeridos'}), 400
        
        placeholders = ','.join(['%s'] * len(lot_nos))
        sql_select = f"""
            SELECT lot_no, wo_id, wo_code, po_code, working_date, line, model_code, 
                   part_no, project, process, plan_count, produced_count, ct, uph, routing, 
                   status, group_no, sequence
            FROM plan_main 
            WHERE lot_no IN ({placeholders})
        """
        planes_originales = execute_query(sql_select, tuple(lot_nos), fetch='all')
        
        if not planes_originales:
            return jsonify({'error': 'No se encontraron planes para reprogramar'}), 404
        
        nuevos_planes_creados = 0
        
        for plan in planes_originales:
            lot_no_original = plan['lot_no']
            plan_count_original = plan['plan_count'] or 0
            produced_count = plan['produced_count'] or 0
            cantidad_pendiente = plan_count_original - produced_count
            
            if cantidad_pendiente <= 0:
                continue
            
            # Generar nuevo lot_no con trazabilidad
            if lot_no_original.count('-') >= 3:
                parts = lot_no_original.rsplit('-', 1)
                lot_no_base = parts[0]
            else:
                lot_no_base = lot_no_original
            
            sql_count = "SELECT COUNT(*) as count FROM plan_main WHERE lot_no LIKE %s AND lot_no <> %s"
            result = execute_query(sql_count, (f"{lot_no_base}-%", lot_no_base), fetch='one')
            count = result['count'] if result else 0
            next_seq = count + 1
            nuevo_lot_no = f"{lot_no_base}-{next_seq:02d}"
            
            sql_insert = """
                INSERT INTO plan_main 
                (lot_no, wo_id, wo_code, po_code, working_date, line, model_code, 
                 part_no, project, process, plan_count, ct, uph, routing, status, 
                 group_no, sequence, created_at)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, NOW())
            """
            
            execute_query(sql_insert, (
                nuevo_lot_no, plan.get('wo_id'), plan.get('wo_code'), plan.get('po_code'),
                new_date, plan.get('line'), plan.get('model_code'), plan.get('part_no'),
                plan.get('project'), plan.get('process'), cantidad_pendiente, plan.get('ct'),
                plan.get('uph'), plan.get('routing'), 'PLAN', plan.get('group_no'), plan.get('sequence')
            ))
            
            nuevos_planes_creados += 1
        
        return jsonify({
            'success': True, 
            'created': nuevos_planes_creados,
            'message': f'{nuevos_planes_creados} nuevo(s) plan(es) creado(s) para {new_date}'
        })
        
    except Exception as e:
        print(f"❌ Error en api_plan_reschedule: {str(e)}")
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500


@produccion_bp.route('/api/plan/export-excel', methods=['POST'])
@login_requerido
def api_plan_export_excel():
    """Exportar planes a Excel"""
    try:
        payload = request.get_json() or {}
        plans = payload.get('plans', [])
        
        if not plans:
            return jsonify({'error': 'No hay datos para exportar'}), 400
        
        wb = Workbook()
        ws = wb.active
        ws.title = 'Plan Producción'
        
        headers = ['Sec', 'LOT NO', 'WO', 'PO', 'Fecha', 'Línea', 'Turno', 'Modelo', 'Part No', 
                   'Proyecto', 'Proceso', 'CT', 'UPH', 'Plan', 'Producido', 'Status', 'Tiempo', 
                   'Inicio', 'Fin', 'Grupo', 'Extra']
        ws.append(headers)
        
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal='center')
        
        for p in plans:
            if p.get('isGroupHeader'):
                ws.append([p.get('groupTitle', f"GRUPO {p.get('groupIndex', 0)+1}")])
                continue
            ws.append([
                p.get('secuencia',''), p.get('lot_no',''), p.get('wo_code',''), p.get('po_code',''),
                p.get('working_date',''), p.get('line',''), p.get('turno',''), p.get('model_code',''),
                p.get('part_no',''), p.get('project',''), p.get('process',''), p.get('ct',''),
                p.get('uph',''), p.get('plan_count',''), p.get('produced',''), p.get('status',''),
                p.get('tiempo_produccion',''), p.get('inicio',''), p.get('fin',''), p.get('grupo',''),
                p.get('extra','')
            ])
        
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        ts = datetime.utcnow().strftime('%Y%m%d_%H%M')
        
        return send_file(bio, as_attachment=True, download_name=f'Plan_Produccion_{ts}.xlsx',
                        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# ============== API RAW ==============

@produccion_bp.route('/api/raw/search', methods=['GET'])
@login_requerido
def api_raw_search():
    """Buscar datos en la tabla RAW por part_no o model"""
    try:
        part_no = request.args.get('part_no', '').strip()
        if not part_no:
            return jsonify({'error': 'part_no requerido'}), 400
        
        sql = """
            SELECT part_no, model, project, c_t as ct, uph 
            FROM raw 
            WHERE TRIM(model) = %s 
               OR TRIM(part_no) = %s 
               OR TRIM(part_no) LIKE %s
               OR UPPER(TRIM(part_no)) = UPPER(%s)
            LIMIT 1
        """
        params = (part_no, part_no, f'%{part_no}%', part_no)
        result = execute_query(sql, params, fetch='all')
        
        if result and isinstance(result, (list, tuple)) and len(result) > 0:
            row = result[0]
            data = {
                'part_no': row.get('part_no', '') or '',
                'model': row.get('model', '') or '',
                'model_code': row.get('model', '') or '',
                'project': row.get('project', '') or '',
                'ct': str(row.get('ct', '0') or '0'),
                'uph': str(row.get('uph', '0') or '0')
            }
            return jsonify([data])
        else:
            return jsonify([])
            
    except Exception as e:
        print(f"Error en api_raw_search: {e}")
        return jsonify({'error': str(e)}), 500


# ============== API PLAN-MAIN LIST ==============

@produccion_bp.route('/api/plan-main/list', methods=['GET'])
@login_requerido
def api_plan_main_list():
    """Listar planes con filtros avanzados"""
    try:
        q = request.args.get('q', '').strip()
        linea = request.args.get('linea')
        desde = request.args.get('desde')
        hasta = request.args.get('hasta')
        solo_pendientes = request.args.get('solo_pendientes') == 'true'
        
        where = []
        params = []
        
        if q:
            where.append('(lot_no LIKE %s OR part_no LIKE %s OR model_code LIKE %s)')
            qv = f"%{q}%"
            params.extend([qv, qv, qv])
        if linea and linea not in ('Todos', 'ALL'):
            where.append('line = %s')
            params.append(linea)
        if desde:
            where.append('DATE(working_date) >= %s')
            params.append(desde)
        if hasta:
            where.append('DATE(working_date) <= %s')
            params.append(hasta)
        if solo_pendientes:
            where.append("status = 'PLAN'")
        
        sql = (
            "SELECT id, lot_no, part_no, model_code, line, working_date, "
            "COALESCE(plan_count,0) AS qty, COALESCE(produced_count,0) AS producido, "
            "GREATEST(COALESCE(plan_count,0)-COALESCE(produced_count,0),0) AS falta, "
            "COALESCE(ct,0) AS ct, COALESCE(uph,0) AS uph, status, process "
            "FROM plan_main"
        )
        
        if where:
            sql += ' WHERE ' + ' AND '.join(where)
        sql += ' ORDER BY working_date DESC, created_at DESC'
        
        rows = execute_query(sql, tuple(params) if params else None, fetch='all')
        
        out = []
        for r in rows:
            qty = r['qty'] if isinstance(r, dict) else r[6]
            producido = r.get('producido', 0) if isinstance(r, dict) else r[7]
            pct = int(round((producido/qty)*100, 0)) if qty else 0
            
            out.append({
                'id': r['id'] if isinstance(r, dict) else r[0],
                'lote': r['lot_no'] if isinstance(r, dict) else r[1],
                'nparte': r['part_no'] if isinstance(r, dict) else r[2],
                'modelo': r['model_code'] if isinstance(r, dict) else r[3],
                'linea': r['line'] if isinstance(r, dict) else r[4],
                'fecha_inicio': str((r['working_date'] if isinstance(r, dict) else r[5]) or '')[:10],
                'qty': qty,
                'producido': producido,
                'falta': max(0, qty - producido),
                'ct': r['ct'] if isinstance(r, dict) else r[9],
                'uph': r['uph'] if isinstance(r, dict) else r[10],
                'estatus': r['status'] if isinstance(r, dict) else r[11],
                'process': r['process'] if isinstance(r, dict) else r[12],
            })
        
        return jsonify(out)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# ============== TEMPLATES ==============

@produccion_bp.route('/cargar_template', methods=['POST'])
@login_requerido
def cargar_template():
    """Cargar un template dinámicamente"""
    template_path = None
    try:
        data = request.get_json()
        template_path = data.get('template_path')
        
        if not template_path:
            return jsonify({'error': 'No se especificó la ruta del template'}), 400
        
        if '..' in template_path or template_path.startswith('/'):
            return jsonify({'error': 'Ruta de template no válida'}), 400
        
        html_content = render_template(template_path)
        return html_content
        
    except Exception as e:
        template_name = template_path if template_path else 'unknown'
        print(f"Error al cargar template {template_name}: {str(e)}")
        return jsonify({'error': f'Error al cargar el template: {str(e)}'}), 500
