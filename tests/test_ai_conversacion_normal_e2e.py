"""Conversacion normal: un usuario sube un Excel ajeno al MES y pregunta.

El caso que el marco viejo del prompt trataba mal: un archivo que no es de
ningun modulo del MES (una lista de retardos de RH) y preguntas de sentido
comun sobre el. El asistente debe contestarlas con su propio conocimiento,
sin negarse, sin redirigir al MES y sin tocar reportes ni escribir nada.
"""

import json

from app.api.portal import ai_assistant, ai_openai
from conftest import _turno

ARCHIVO = "Retardos Agosto 2026.xlsx"

# Nombres y minutos inventados para la prueba.
RETARDOS = [
    ("Empleado", "Fecha", "Hora entrada", "Minutos retardo"),
    ("Ana Ramírez", "2026-08-03", "07:42", 12),
    ("Luis Ortega", "2026-08-03", "08:15", 45),
    ("Ana Ramírez", "2026-08-05", "07:51", 21),
    ("María Fuentes", "2026-08-06", "07:33", 3),
    ("Luis Ortega", "2026-08-07", "08:02", 32),
    ("Luis Ortega", "2026-08-11", "07:58", 28),
    ("Ana Ramírez", "2026-08-12", "07:44", 14),
]


def _subir_excel(mundo, conversacion=7, ref="ret001"):
    """Deja el xlsx donde el blueprint lo va a encontrar y devuelve su ref."""
    from openpyxl import Workbook

    carpeta = ai_assistant._upload_root() / str(conversacion)
    carpeta.mkdir(parents=True, exist_ok=True)
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Retardos"
    for fila in RETARDOS:
        hoja.append(list(fila))
    libro.save(carpeta / (ref + ".xlsx"))
    (carpeta / (ref + ".name")).write_text(ARCHIVO, "utf-8")
    return ref


def test_pregunta_normal_sobre_un_excel_ajeno_al_mes(client, mundo):
    """"dime las personas con mas retardos" sobre un archivo que no es del MES."""
    ref = _subir_excel(mundo)
    mundo.guion = [[
        "Luis Ortega es quien mas acumula: 3 retardos y 105 minutos. "
        "Le sigue Ana Ramírez con 3 retardos y 47 minutos, y "
        "María Fuentes con 1 retardo de 3 minutos."
    ]]

    cuerpo = _turno(client, "dime las personas con mas retardos", file_refs=[ref])

    # El contenido del Excel llego al modelo, con nombres y minutos.
    visto = json.dumps(mundo.vistas_por_el_modelo, ensure_ascii=False, default=str)
    assert "Luis Ortega" in visto
    assert "Minutos retardo" in visto
    assert ARCHIVO in visto

    # Contesto de corrido, sin llamar reportes del MES y sin escribir.
    assert "Luis Ortega" in cuerpo
    assert mundo.ejecuciones == [], "una pregunta sobre el adjunto no usa tools del MES"
    assert mundo.escrituras == []


def test_el_prompt_no_le_manda_negarse_ni_redirigir_al_mes(mundo):
    """El marco nuevo: contesta temas ajenos al MES sin disclaimers ni rebote."""
    info = ai_assistant._uploaded_file_info(7, _subir_excel(mundo))
    prompt = ai_openai.build_instructions({"language": "es", "attachment": info})

    assert info["kind"] == "excel"
    assert ARCHIVO in prompt
    assert "El archivo ya llegó al servidor" in prompt
    assert "No digas que no fue recibido" in prompt
    # Marco de capacidad, no de rebote al MES.
    assert "Conversas como cualquier asistente capaz" in prompt
    assert "sin redirigir al usuario al MES" in prompt
    assert "asistente oficial de solo lectura" not in prompt


def test_en_el_seguimiento_persiste_la_respuesta_previa_no_las_filas(client, mundo):
    """Comportamiento real, util de conocer: las filas del Excel viajan SOLO en
    el turno que trajo el archivo (_turn_attachment_parts). En los turnos
    siguientes el modelo se apoya en su propia respuesta anterior, que si queda
    en el historial. Para volver a leer el detalle hay que readjuntar."""
    ref = _subir_excel(mundo)
    primera = "Luis Ortega acumula mas: 3 retardos, 105 minutos."
    mundo.guion = [[primera]]
    _turno(client, "quien acumula mas minutos?", file_refs=[ref])

    primer_input = json.dumps(mundo.vistas_por_el_modelo[0], ensure_ascii=False, default=str)
    assert "Minutos retardo" in primer_input, "las filas deben llegar en SU turno"

    for pregunta, respuesta in [
        ("y cuantos empleados distintos hay?", "Tres: Ana, Luis y Maria."),
        ("cual fue el retardo mas grande?", "45 minutos, Luis el 3 de agosto."),
        ("hazme un promedio por persona", "Luis 35, Ana 15.7, Maria 3."),
    ]:
        mundo.guion = [[respuesta]]
        cuerpo = _turno(client, pregunta)
        assert respuesta[:12] in cuerpo

    ultimo = json.dumps(mundo.vistas_por_el_modelo[-1], ensure_ascii=False, default=str)
    assert "Minutos retardo" not in ultimo, "las filas crudas no se reenvian cada turno"
    assert primera in ultimo, "pero si conserva su conclusion anterior"
    assert mundo.ejecuciones == [], "ninguna de estas preguntas necesita el MES"
    assert mundo.escrituras == []


def test_pedir_el_resultado_en_excel_ofrece_la_tool_no_un_csv_de_texto(client, mundo):
    """Exportar lo que salio de un adjunto: excel_desde_tabla, no "no puedo"."""
    ref = _subir_excel(mundo)
    mundo.guion = [["Luis Ortega encabeza con 105 minutos."]]
    _turno(client, "dime las personas con mas retardos", file_refs=[ref])

    prompt = ai_openai.build_instructions({"language": "es", "table_excel_enabled": True})
    assert "excel_desde_tabla" in prompt
    assert "no puedes generar Excel" in prompt
    assert "ofrezcas el texto en CSV" in prompt
    assert "powerpoint_desde_tabla" in prompt


# ---------------------------------------------------------------------------
# Plantilla estilo "SOLICITUD DE TIEMPO EXTRA": titulo suelto arriba, encabezado
# en la fila 6, columnas vacias intercaladas y varias hojas por dia.
# ---------------------------------------------------------------------------
ENCABEZADO_TE = ["NO.", "AREA", "PUESTO", "System Nombre", "ACTIVIDAD",
                 "TIEMPO", "HRS", "NO. EMPLEADO"]
TURNOS_TE = [
    (1, "Produccion", "M3/M4", "FLORES MUNIZ VERONICA", "LIDER", "17:30", 4, 429),
    (2, "Produccion", "Coating", "RAMIREZ RAMIREZ VIANEEY", "M4 P1", "17:30", 4, 697),
    (3, "Produccion", "Harness/D3", "FUENTES GONZALEZ ERIKA", "M4 P2", "17:30", 2, 2105),
]


def _plantilla_tiempo_extra(ref="te001", conversacion=7):
    """Mismo esqueleto que la plantilla real: B como columna vacia de margen."""
    from openpyxl import Workbook

    carpeta = ai_assistant._upload_root() / str(conversacion)
    carpeta.mkdir(parents=True, exist_ok=True)
    libro = Workbook()
    for dia, filas in [("LUNES", TURNOS_TE), ("MARTES", TURNOS_TE[:2])]:
        hoja = libro.create_sheet(dia)
        hoja["C2"] = "SOLICITUD DE TIEMPO EXTRA -"
        hoja["K3"] = "Fecha"
        hoja["K4"] = "2026-09-07"
        for col, titulo in enumerate(ENCABEZADO_TE, start=3):   # arranca en C
            hoja.cell(row=6, column=col, value=titulo)
        for i, fila in enumerate(filas):
            for col, valor in enumerate(fila, start=3):
                hoja.cell(row=7 + i, column=col, value=valor)
    del libro["Sheet"]
    libro.save(carpeta / (ref + ".xlsx"))
    (carpeta / (ref + ".name")).write_text("PLANTILLA DE TIEMPO EXTRA.xlsx", "utf-8")
    return ref


def test_plantilla_con_encabezado_abajo_llega_legible_y_sin_columnas_vacias(mundo):
    """El volcado conserva filas/letras reales pero tira las columnas vacias."""
    ref = _plantilla_tiempo_extra()
    ruta = ai_assistant._upload_root() / "7" / (ref + ".xlsx")
    texto = ai_assistant._excel_a_texto(ruta.read_bytes())

    # Las dos hojas de dia, con sus datos.
    assert "# Hoja: LUNES" in texto and "# Hoja: MARTES" in texto
    assert "FLORES MUNIZ VERONICA" in texto
    assert "NO. EMPLEADO" in texto

    # La columna A y la B (margen) no gastan espacio, pero las letras reales
    # se conservan para poder pedir una edicion por celda.
    cabecera = [l for l in texto.splitlines() if l.startswith("fila | ")][0]
    assert cabecera.startswith("fila | C | ")
    assert " | A | " not in cabecera

    # El numero de fila real sobrevive: el encabezado sigue siendo la fila 6.
    fila6 = [l for l in texto.splitlines() if l.startswith("6 | ")][0]
    assert "NO. EMPLEADO" in fila6


def test_el_prompt_avisa_que_el_encabezado_puede_no_estar_en_la_fila_1(mundo):
    info = ai_assistant._uploaded_file_info(7, _plantilla_tiempo_extra())
    prompt = ai_openai.build_instructions({"language": "es", "attachment": info})
    assert "el encabezado casi nunca está en la fila 1" in prompt
    assert "localiza la hoja relevante" in prompt
    assert "las letras pueden saltarse" in prompt


def test_pregunta_de_negocio_sobre_la_plantilla_de_tiempo_extra(client, mundo):
    """"quien hizo mas horas extra" sobre la plantilla, sin tocar el MES."""
    ref = _plantilla_tiempo_extra()
    mundo.guion = [[
        "En LUNES y MARTES, Veronica Flores y Vianeey Ramirez acumulan 8 h cada "
        "una (4 h por dia) y Erika Fuentes 2 h. Encabezado en la fila 6."
    ]]
    cuerpo = _turno(client, "quien hizo mas horas extra?", file_refs=[ref])

    visto = json.dumps(mundo.vistas_por_el_modelo, ensure_ascii=False, default=str)
    assert "FLORES MUNIZ VERONICA" in visto
    assert "NO. EMPLEADO" in visto
    assert "Veronica Flores" in cuerpo
    assert mundo.ejecuciones == [], "una plantilla ajena al MES no usa reportes"
    assert mundo.escrituras == []


# ---------------------------------------------------------------------------
# Libro grande: viaja como indice y el modelo pide la hoja que necesita.
# ---------------------------------------------------------------------------
def _libro_grande(ref="big001", conversacion=7, hojas=("ASISTENCIA", "LUNES", "MARTES")):
    """Supera _EXCEL_INDICE_DESDE: una hoja padron ancha y dos de tiempo extra."""
    from openpyxl import Workbook

    carpeta = ai_assistant._upload_root() / str(conversacion)
    carpeta.mkdir(parents=True, exist_ok=True)
    libro = Workbook()
    for nombre in hojas:
        hoja = libro.create_sheet(nombre)
        if nombre == "ASISTENCIA":
            hoja.append(["No.", "System Nombre", "Escolaridad", "Sexo", "Fecha de Ingreso"])
            for i in range(400):
                hoja.append([i, f"EMPLEADO NUMERO {i} APELLIDO", "UNIVERSIDAD", "F", "2017-03-29"])
        else:
            for col, titulo in enumerate(ENCABEZADO_TE, start=3):
                hoja.cell(row=6, column=col, value=titulo)
            for i, fila in enumerate(TURNOS_TE):
                for col, valor in enumerate(fila, start=3):
                    hoja.cell(row=7 + i, column=col, value=valor)
    del libro["Sheet"]
    libro.save(carpeta / (ref + ".xlsx"))
    (carpeta / (ref + ".name")).write_text("LIBRO GRANDE.xlsx", "utf-8")
    return ref


def test_libro_grande_viaja_como_mapa_no_como_volcado(mundo):
    ref = _libro_grande()
    info = ai_assistant._uploaded_file_info(7, ref)
    partes = ai_assistant._attachment_input_parts(7, ref, info)
    texto = partes[0]["text"]

    # Llega el mapa con dimensiones REALES, no las filas.
    assert "MAPA del Excel" in texto
    assert "excel_leer_hoja" in texto and "excel_contar_en_rango" in texto
    assert "401 filas x 5 columnas" in texto, texto
    assert "encabezados en la fila 1" in texto
    assert "A=No., B=System Nombre" in texto
    assert "EMPLEADO NUMERO 300 APELLIDO" not in texto, "el detalle NO debe viajar"

    volcado = ai_assistant._excel_a_texto(
        (ai_assistant._upload_root() / "7" / (ref + ".xlsx")).read_bytes())
    assert len(texto) < len(volcado) / 5


def test_libro_chico_sigue_viajando_completo(mundo):
    """El umbral no debe penalizar el caso facil: una plantilla chica va entera."""
    ref = _plantilla_tiempo_extra(ref="chico01")
    partes = ai_assistant._attachment_input_parts(7, ref, ai_assistant._uploaded_file_info(7, ref))
    texto = partes[0]["text"]

    assert "FLORES MUNIZ VERONICA" in texto, "un libro chico va con su detalle"
    assert "excel_leer_hoja" not in texto


def test_el_modelo_pide_una_hoja_y_recibe_solo_esa(client, mundo):
    """Ciclo completo: indice -> excel_leer_hoja(LUNES) -> respuesta."""
    ref = _libro_grande()
    mundo.guion = [[
        {"tool": "excel_leer_hoja", "args": {"hoja": "LUNES"}},
        "El lunes Veronica Flores hizo 4 h extra, la mas alta.",
    ]]
    cuerpo = _turno(client, "quien hizo mas horas extra el lunes?", file_refs=[ref])

    # La tool corrio y devolvio SOLO esa hoja.
    assert [e["tool_name"] for e in mundo.ejecuciones] == ["excel_leer_hoja"]
    assert mundo.ejecuciones[0]["result"]["hoja"] == "LUNES"

    entregado = json.dumps(mundo.vistas_por_el_modelo[-1], ensure_ascii=False, default=str)
    assert "FLORES MUNIZ VERONICA" in entregado, "la hoja pedida si llega"
    assert "EMPLEADO NUMERO 300" not in entregado, "ASISTENCIA no se cuela"
    assert "Veronica Flores" in cuerpo
    assert mundo.escrituras == []


def test_pedir_una_hoja_inexistente_devuelve_los_nombres_reales(client, mundo):
    """No adivinar: se le dicen las hojas disponibles para que reintente."""
    ref = _libro_grande()
    mundo.guion = [[
        {"tool": "excel_leer_hoja", "args": {"hoja": "DOMINGO"}},
        "Ese libro no tiene hoja DOMINGO; tiene ASISTENCIA, LUNES y MARTES.",
    ]]
    _turno(client, "que dice la hoja del domingo?", file_refs=[ref])

    entregado = json.dumps(mundo.vistas_por_el_modelo[-1], ensure_ascii=False, default=str)
    assert "no existe en el libro" in entregado
    assert "LUNES" in entregado and "MARTES" in entregado


# ---------------------------------------------------------------------------
# Matriz ancha: personas en filas, un dia por columna. La forma que no cabia.
# ---------------------------------------------------------------------------
CODIGOS = {0: "R", 1: "A", 2: "A", 3: "R", 4: "A"}


def _matriz_asistencia(ref="mat001", conversacion=7, personas=60, dias=120):
    """Encabezado en la fila 3, atributos en A-D y un dia por columna desde E."""
    from datetime import date, timedelta

    from openpyxl import Workbook

    carpeta = ai_assistant._upload_root() / str(conversacion)
    carpeta.mkdir(parents=True, exist_ok=True)
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Activos"
    hoja["B1"] = "Asistencia"
    for col, titulo in enumerate(["No.", "System Nombre", "Area", "Turno"], start=1):
        hoja.cell(row=3, column=col, value=titulo)
    inicio = date(2026, 1, 1)
    for d in range(dias):
        hoja.cell(row=3, column=5 + d, value=inicio + timedelta(days=d))
    for i in range(personas):
        fila = 4 + i
        hoja.cell(row=fila, column=1, value=100 + i)
        hoja.cell(row=fila, column=2, value=f"PERSONA {i:03d} APELLIDO")
        hoja.cell(row=fila, column=3, value="Produccion")
        hoja.cell(row=fila, column=4, value="DIA")
        # La persona i tiene exactamente i retardos, para poder verificar.
        for d in range(dias):
            hoja.cell(row=fila, column=5 + d, value="R" if d < i else "A")
    libro.save(carpeta / (ref + ".xlsx"))
    (carpeta / (ref + ".name")).write_text("Libro asistencia.xlsx", "utf-8")
    return ref


def test_el_mapa_ubica_las_columnas_de_fecha_por_anio(mundo):
    """Sin esto el modelo no sabe donde empieza 2026 y responde sobre otro rango."""
    ref = _matriz_asistencia()
    data = (ai_assistant._upload_root() / "7" / (ref + ".xlsx")).read_bytes()
    mapa = ai_assistant._excel_indice(data, "Libro asistencia.xlsx")

    assert "encabezados en la fila 3" in mapa
    assert "B=System Nombre" in mapa
    assert "Columnas de fecha: E=2026-01-01" in mapa
    assert "120 columnas" in mapa
    assert "2026=E.." in mapa


def test_contar_recorre_la_hoja_completa_no_un_fragmento(mundo):
    """El conteo es exacto sobre TODAS las filas y columnas, sin truncar."""
    ref = _matriz_asistencia(personas=60, dias=120)
    data = (ai_assistant._upload_root() / "7" / (ref + ".xlsx")).read_bytes()
    info = next(h for h in ai_assistant._excel_mapa(data) if h["hoja"] == "Activos")

    r = ai_assistant._excel_contar_rango(data, "Activos", info, {
        "columna_etiqueta": "B", "desde_columna": "E", "hasta_columna": "DZ",
        "valores": ["R"], "top": 3,
    })

    # La persona i tiene i retardos: la 59 encabeza con 59.
    assert r["ranking"][0] == {"etiqueta": "PERSONA 059 APELLIDO", "conteo": 59}
    assert r["ranking"][1]["conteo"] == 58
    assert r["filas_con_coincidencias"] == 59, "la persona 000 no tiene retardos"
    assert r["total_coincidencias"] == sum(range(60))
    # Filas mas alla del vistazo automatico tambien cuentan.
    assert info["filas"] > 60 or info["columnas"] > 40


def test_contar_no_distingue_mayusculas_y_exige_un_valor(mundo):
    ref = _matriz_asistencia(personas=5, dias=10)
    data = (ai_assistant._upload_root() / "7" / (ref + ".xlsx")).read_bytes()
    info = next(h for h in ai_assistant._excel_mapa(data) if h["hoja"] == "Activos")

    minus = ai_assistant._excel_contar_rango(data, "Activos", info, {
        "columna_etiqueta": "B", "desde_columna": "E", "hasta_columna": "N",
        "valores": ["r"], "top": 5,
    })
    assert minus["total_coincidencias"] == sum(range(5))

    try:
        ai_assistant._excel_contar_rango(data, "Activos", info, {
            "columna_etiqueta": "B", "desde_columna": "E", "hasta_columna": "N",
            "valores": [], "top": 5,
        })
        raise AssertionError("acepto contar sin valores")
    except ValueError as exc:
        assert "al menos un valor" in str(exc)


def test_leer_un_rango_grande_sugiere_contar_en_vez_de_truncar(mundo):
    """No se trunca en silencio: se dice que pida menos o que use el conteo."""
    ref = _matriz_asistencia(personas=60, dias=120)
    data = (ai_assistant._upload_root() / "7" / (ref + ".xlsx")).read_bytes()
    info = next(h for h in ai_assistant._excel_mapa(data) if h["hoja"] == "Activos")
    info = {**info, "filas": 50000, "columnas": 500}

    try:
        ai_assistant._excel_leer_rango(data, "Activos", info, {
            "desde_fila": 1, "hasta_fila": 50000,
            "desde_columna": "A", "hasta_columna": "SF",
        })
        raise AssertionError("acepto un rango imposible")
    except ValueError as exc:
        assert "excel_contar_en_rango" in str(exc)


def test_ciclo_completo_quien_tiene_mas_retardos(client, mundo):
    """La pregunta que fallaba: mapa -> contar -> respuesta, sin readjuntar."""
    ref = _matriz_asistencia(personas=60, dias=120)
    mundo.guion = [[
        {"tool": "excel_contar_en_rango",
         "args": {"hoja": "Activos", "columna_etiqueta": "B", "desde_columna": "E",
                  "hasta_columna": "DZ", "valores": ["R"], "top": 5}},
        "PERSONA 059 APELLIDO encabeza con 59 retardos en 2026.",
    ]]
    cuerpo = _turno(client, "dime las personas que mas retardos tienen", file_refs=[ref])

    assert [e["tool_name"] for e in mundo.ejecuciones] == ["excel_contar_en_rango"]
    entregado = json.dumps(mundo.vistas_por_el_modelo[-1], ensure_ascii=False, default=str)
    assert "PERSONA 059 APELLIDO" in entregado
    assert "PERSONA 059 APELLIDO" in cuerpo
    assert mundo.escrituras == []


def test_el_prompt_prohibe_decir_que_quedo_truncado(mundo):
    ref = _matriz_asistencia()
    info = ai_assistant._uploaded_file_info(7, ref)
    prompt = ai_openai.build_instructions({"language": "es", "attachment": info})
    assert "ninguna celda esta fuera de alcance" in prompt
    assert "excel_contar_en_rango" in prompt
    assert "jamas cuentes a mano sobre un fragmento visible" in prompt


def test_el_libro_sigue_disponible_en_turnos_posteriores(client, mundo):
    """El caso "ahora con mas faltas": antes moria pidiendo readjuntar."""
    ref = _matriz_asistencia(personas=40, dias=90)

    mundo.guion = [[
        {"tool": "excel_contar_en_rango",
         "args": {"hoja": "Activos", "columna_etiqueta": "B", "desde_columna": "E",
                  "hasta_columna": "CT", "valores": ["R"], "top": 5}},
        "PERSONA 039 encabeza con 39 retardos.",
    ]]
    _turno(client, "dime las personas con mas retardos", file_refs=[ref])

    # Turno siguiente SIN adjuntar: la tool debe seguir funcionando.
    mundo.guion = [[
        {"tool": "excel_contar_en_rango",
         "args": {"hoja": "Activos", "columna_etiqueta": "B", "desde_columna": "E",
                  "hasta_columna": "CT", "valores": ["A"], "top": 5}},
        "Ahora por faltas: PERSONA 000 encabeza.",
    ]]
    cuerpo = _turno(client, "ahora con mas faltas")

    ejecutadas = [e["tool_name"] for e in mundo.ejecuciones]
    assert ejecutadas == ["excel_contar_en_rango", "excel_contar_en_rango"], ejecutadas
    assert "PERSONA 000" in cuerpo
    # Y el resultado dice sobre que archivo trabajo.
    assert mundo.ejecuciones[-1]["result"] is not None
    entregado = json.dumps(mundo.vistas_por_el_modelo[-1], ensure_ascii=False, default=str)
    assert "Libro asistencia.xlsx" in entregado


def test_el_prompt_avisa_que_el_libro_anterior_sigue_vivo(client, mundo):
    ref = _matriz_asistencia(personas=10, dias=30)
    mundo.guion = [["Listo."]]
    _turno(client, "resume el archivo", file_refs=[ref])

    prompt = ai_openai.build_instructions({
        "language": "es", "libro_vigente": "Libro asistencia.xlsx",
    })
    assert "SIGUE" in prompt and "disponible en el servidor" in prompt
    assert "Nunca pidas que lo readjunten" in prompt


def test_el_excel_generado_reporta_sus_filas_no_cero(client, mundo, monkeypatch):
    """La tarjeta decia "0 filas" porque row_count iba hardcodeado a 0."""
    registrados = []

    def _fake_register(**kwargs):
        registrados.append(kwargs)
        return {"id": "art-1", "type": "xlsx", "title": kwargs.get("title"),
                "filename": kwargs.get("filename"), "row_count": kwargs.get("row_count")}

    monkeypatch.setattr(ai_assistant, "register_file_artifact", _fake_register)
    monkeypatch.setattr(ai_assistant, "check_quota",
                        lambda *_a, **_k: (True, None, {}, {}))

    ref = _matriz_asistencia(personas=5, dias=10)
    mundo.guion = [[
        {"tool": "excel_desde_tabla",
         "args": {"titulo": "Top retardos", "columnas": ["Persona", "Retardos"],
                  "filas": [["PERSONA 004", 4], ["PERSONA 003", 3], ["PERSONA 002", 2]]}},
        "Listo, adjunte el Excel.",
    ]]
    _turno(client, "pasalo a excel", file_refs=[ref])

    assert registrados, "no se registro el artefacto"
    assert registrados[0]["row_count"] == 3, registrados[0]


def test_agotar_el_presupuesto_de_tools_cierra_la_respuesta_no_la_mata(client, mundo, monkeypatch):
    """Antes: "Se alcanzo el limite de herramientas" y se perdia todo el trabajo."""
    monkeypatch.setenv("AI_MAX_TOOL_CALLS", "2")
    ref = _matriz_asistencia(personas=40, dias=90)

    # El modelo insiste en llamar tools mas alla del presupuesto.
    llamada = {"tool": "excel_contar_en_rango",
               "args": {"hoja": "Activos", "columna_etiqueta": "B", "desde_columna": "E",
                        "hasta_columna": "CT", "valores": ["R"], "top": 3}}
    # Cuatro llamadas en UNA sola vuelta: mas de las que caben en el presupuesto.
    mundo.guion = [[[llamada, llamada, llamada, llamada],
                    "Con lo que alcance a revisar: PERSONA 039 encabeza. Quedo pendiente el resto."]]

    cuerpo = _turno(client, "dame el detalle completo", file_refs=[ref])

    assert "Se alcanz" not in cuerpo, "no debe salir el error duro"
    assert "PERSONA 039" in cuerpo, "debe responder con lo que reunio"
    # Solo se ejecutaron las tools del presupuesto.
    assert len([e for e in mundo.ejecuciones if e["status"] == "success"]) <= 2
    # Y al modelo se le dijo que cerrara.
    entregado = json.dumps(mundo.vistas_por_el_modelo[-1], ensure_ascii=False, default=str)
    assert "No llames mas herramientas" in entregado


def test_el_mapa_viaja_en_turnos_posteriores_para_no_preguntar_columnas(client, mundo):
    """Antes preguntaba "en que columna esta la fecha de ingreso?" teniendo el mapa."""
    ref = _matriz_asistencia(personas=40, dias=90)
    mundo.guion = [["Listo."]]
    _turno(client, "dime los retardos", file_refs=[ref])

    mundo.guion = [["Uso la columna B y las fechas desde E."]]
    _turno(client, "ahora solo los de cierta antiguedad")

    # El mapa viaja en las instrucciones, no en los mensajes.
    prompt = mundo.prompts[-1]
    assert "MAPA del Excel" in prompt, "el mapa debe seguir presente"
    assert "B=System Nombre" in prompt, "y decir que hay en cada columna"
    assert "encabezados en la fila 3" in prompt


def _matriz_con_ingreso(ref="ing001", conversacion=7, personas=30, dias=60):
    """Como _matriz_asistencia pero con fecha de ingreso en la columna C."""
    from datetime import date, timedelta

    from openpyxl import Workbook

    carpeta = ai_assistant._upload_root() / str(conversacion)
    carpeta.mkdir(parents=True, exist_ok=True)
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Activos"
    for col, titulo in enumerate(["No.", "System Nombre", "Fecha de Ingreso"], start=1):
        hoja.cell(row=3, column=col, value=titulo)
    inicio = date(2026, 1, 1)
    for d in range(dias):
        hoja.cell(row=3, column=4 + d, value=inicio + timedelta(days=d))
    for i in range(personas):
        fila = 4 + i
        hoja.cell(row=fila, column=1, value=100 + i)
        hoja.cell(row=fila, column=2, value=f"PERSONA {i:03d}")
        # La mitad son nuevos (ingreso 2026), la mitad antiguos (2020).
        hoja.cell(row=fila, column=3, value=date(2026, 3, 1) if i % 2 else date(2020, 5, 5))
        for d in range(dias):
            hoja.cell(row=fila, column=4 + d, value="R" if d < i else "A")
    libro.save(carpeta / (ref + ".xlsx"))
    (carpeta / (ref + ".name")).write_text("Asistencia con ingreso.xlsx", "utf-8")
    return ref


def _contar(data, **args):
    info = next(h for h in ai_assistant._excel_mapa(data) if h["hoja"] == "Activos")
    base = {"columna_etiqueta": "B", "desde_columna": "D", "hasta_columna": "BK"}
    return ai_assistant._excel_contar_rango(data, "Activos", info, {**base, **args})


def test_orden_ascendente_incluye_a_los_que_tienen_cero(mundo):
    """Antes se descartaban las filas con cero, que son justo la respuesta."""
    ref = _matriz_con_ingreso(personas=30, dias=60)
    data = (ai_assistant._upload_root() / "7" / (ref + ".xlsx")).read_bytes()

    menos = _contar(data, valores=["R"], orden="asc", top=3)
    assert menos["ranking"][0] == {"etiqueta": "PERSONA 000", "conteo": 0}
    assert menos["filas_evaluadas"] == 30, "todas las personas se evaluan"
    assert menos["filas_con_coincidencias"] == 29, "solo la 000 tiene cero"

    mas = _contar(data, valores=["R"], orden="desc", top=3)
    assert mas["ranking"][0] == {"etiqueta": "PERSONA 029", "conteo": 29}


def test_filtrar_por_fecha_de_ingreso_acota_sin_leer_filas(mundo):
    """"solo los de mas de un anio" en la misma llamada, sin lecturas extra."""
    ref = _matriz_con_ingreso(personas=30, dias=60)
    data = (ai_assistant._upload_root() / "7" / (ref + ".xlsx")).read_bytes()

    antiguos = _contar(data, valores=["R"], orden="desc", top=50,
                       filtro_columna="C", filtro_hasta="2025-01-01")

    # Solo los pares (ingreso 2020) pasan el filtro.
    assert antiguos["filas_evaluadas"] == 15
    assert antiguos["filas_descartadas_por_filtro"] == 15
    assert all(int(f["etiqueta"].split()[-1]) % 2 == 0 for f in antiguos["ranking"])
    assert antiguos["ranking"][0]["etiqueta"] == "PERSONA 028"


def test_avisa_cuantos_estan_empatados_en_el_primer_lugar(mundo):
    """Con 2000 personas en cero, elegir 15 es arbitrario: hay que decirlo."""
    ref = _matriz_con_ingreso(personas=30, dias=60)
    data = (ai_assistant._upload_root() / "7" / (ref + ".xlsx")).read_bytes()

    # Nadie tiene el codigo Z: todos empatan en cero.
    r = _contar(data, valores=["Z"], orden="asc", top=5)
    assert r["empatados_en_el_primer_lugar"] == 30
    assert r["total_coincidencias"] == 0


def test_filtro_sin_rango_es_un_error_util(mundo):
    ref = _matriz_con_ingreso(personas=5, dias=10)
    data = (ai_assistant._upload_root() / "7" / (ref + ".xlsx")).read_bytes()
    try:
        _contar(data, valores=["R"], filtro_columna="C")
        raise AssertionError("acepto un filtro sin rango")
    except ValueError as exc:
        assert "filtro_desde" in str(exc)


# ---------------------------------------------------------------------------
# PDF grande: no cabe nativo (mas de 100 paginas), va indexado.
# ---------------------------------------------------------------------------
def _manual_pdf(ref="man001", conversacion=7, paginas=120):
    from pdf_minimo import pdf_con_paginas

    hojas = []
    for i in range(paginas):
        if i == 1:
            hojas.append(["INDICE", "Relay Test ------ 40", "Short Open ------ 55"])
        elif i == 39:
            hojas.append(["Relay Test", "El Relay Test verifica cada relay board.",
                          "Relay board 1 controla los pines 1-128.", "Relay relay relay."])
        elif i == 54:
            hojas.append(["SHORT y OPEN", "SHORT significa continuidad indebida."])
        else:
            hojas.append([f"Pagina {i + 1}", "Texto de relleno del manual."])
    carpeta = ai_assistant._upload_root() / str(conversacion)
    carpeta.mkdir(parents=True, exist_ok=True)
    (carpeta / (ref + ".pdf")).write_bytes(pdf_con_paginas(hojas))
    (carpeta / (ref + ".name")).write_text("AT-01_MANUAL ENG.pdf", "utf-8")
    return ref


def test_un_pdf_largo_viaja_como_mapa_y_no_como_no_disponible(mundo):
    """Antes: 183 paginas -> "[supera el limite]" y el modelo no veia nada."""
    ref = _manual_pdf(paginas=120)
    info = ai_assistant._uploaded_file_info(7, ref)
    partes = ai_assistant._attachment_input_parts(7, ref, info)

    assert len(partes) == 1 and partes[0]["type"] == "input_text"
    texto = partes[0]["text"]
    assert "MAPA del PDF" in texto
    assert "120 paginas" in texto
    assert "pdf_buscar" in texto and "pdf_leer_paginas" in texto
    assert "supera el limite" not in texto
    # El mapa ubica secciones sin traer el manual entero.
    assert "Muestreo de paginas" in texto
    assert len(texto) < 6000, "el mapa debe ser barato"


def test_un_pdf_chico_sigue_viajando_nativo(mundo):
    """El modelo lee mejor un PDF chico con sus diagramas: no se indexa."""
    ref = _manual_pdf(ref="chico", paginas=5)
    partes = ai_assistant._attachment_input_parts(
        7, ref, ai_assistant._uploaded_file_info(7, ref))
    assert partes[0]["type"] == "input_file"
    assert partes[0]["file_data"].startswith("data:application/pdf;base64,")


def test_buscar_en_el_pdf_recorre_todas_las_paginas(mundo):
    ref = _manual_pdf(paginas=120)
    ruta = ai_assistant._upload_root() / "7" / (ref + ".pdf")

    res = ai_assistant._pdf_buscar(ruta, {"texto": "Relay", "max_resultados": 3})
    assert res["paginas_totales"] == 120
    # La pagina que explica el tema va primero; el indice queda al final.
    assert res["resultados"][0]["pagina"] == 40
    assert not res["resultados"][0]["parece_indice"]
    assert "Relay board 1" in res["resultados"][0]["fragmento"]


def test_el_indice_del_manual_no_encabeza_los_resultados(mundo):
    """Una pagina de indice menciona el tema pero no lo explica."""
    ref = _manual_pdf(paginas=120)
    ruta = ai_assistant._upload_root() / "7" / (ref + ".pdf")

    res = ai_assistant._pdf_buscar(ruta, {"texto": "Short", "max_resultados": 5})
    paginas = [h["pagina"] for h in res["resultados"]]
    assert paginas[0] == 55, paginas
    assert 2 in paginas, "el indice sale, pero despues"
    assert paginas.index(55) < paginas.index(2)


def test_leer_paginas_devuelve_el_texto_con_su_numero(mundo):
    ref = _manual_pdf(paginas=120)
    ruta = ai_assistant._upload_root() / "7" / (ref + ".pdf")

    res = ai_assistant._pdf_leer(ruta, {"desde_pagina": 40, "hasta_pagina": 41})
    assert "--- pagina 40 ---" in res["contenido"]
    assert "El Relay Test verifica cada relay board." in res["contenido"]
    assert res["paginas_totales"] == 120


def test_ciclo_completo_que_dice_el_manual(client, mundo):
    """La pregunta que fallaba: buscar en el manual y citar la pagina."""
    ref = _manual_pdf(paginas=120)
    mundo.guion = [[
        {"tool": "pdf_buscar", "args": {"texto": "Relay Test", "max_resultados": 3}},
        "Segun la pagina 40, el Relay Test verifica cada relay board.",
    ]]
    cuerpo = _turno(client, "que dice el manual del Relay Test?", file_refs=[ref])

    assert [e["tool_name"] for e in mundo.ejecuciones] == ["pdf_buscar"]
    entregado = json.dumps(mundo.vistas_por_el_modelo[-1], ensure_ascii=False, default=str)
    assert "Relay board 1" in entregado
    assert "pagina 40" in cuerpo


def test_el_pdf_sigue_disponible_en_turnos_posteriores(client, mundo):
    ref = _manual_pdf(paginas=120)
    mundo.guion = [["Listo."]]
    _turno(client, "que dice el manual?", file_refs=[ref])

    mundo.guion = [[
        {"tool": "pdf_buscar", "args": {"texto": "SHORT", "max_resultados": 2}},
        "SHORT significa continuidad indebida, pagina 55.",
    ]]
    cuerpo = _turno(client, "y que significa SHORT?")

    assert "pagina 55" in cuerpo
    assert mundo.prompts[-1].count("pdf_buscar") >= 1
    assert "Nunca digas que el contenido no quedo disponible" in mundo.prompts[-1]


# ---------------------------------------------------------------------------
# PPTX: texto, tablas y notas del presentador.
# ---------------------------------------------------------------------------
def _presentacion(ref="ppt001", conversacion=7):
    from pptx import Presentation
    from pptx.util import Inches

    carpeta = ai_assistant._upload_root() / str(conversacion)
    carpeta.mkdir(parents=True, exist_ok=True)
    pres = Presentation()

    portada = pres.slides.add_slide(pres.slide_layouts[0])
    portada.shapes.title.text = "Resultados OQC Agosto 2026"
    portada.placeholders[1].text = "Planta ILSAN"

    detalle = pres.slides.add_slide(pres.slide_layouts[5])
    detalle.shapes.title.text = "Defectos por linea"
    tabla = detalle.shapes.add_table(
        3, 2, Inches(1), Inches(2), Inches(4), Inches(1.5)).table
    tabla.cell(0, 0).text = "Linea"
    tabla.cell(0, 1).text = "Defectos"
    tabla.cell(1, 0).text = "M1"
    tabla.cell(1, 1).text = "37"
    tabla.cell(2, 0).text = "M4"
    tabla.cell(2, 1).text = "12"
    detalle.notes_slide.notes_text_frame.text = "M1 subio por el cambio de molde."

    pres.save(carpeta / (ref + ".pptx"))
    (carpeta / (ref + ".name")).write_text("Resultados OQC.pptx", "utf-8")
    return ref


def test_pptx_es_un_adjunto_aceptado(mundo):
    """Antes ni se podia subir: la extension no estaba en el catalogo."""
    assert ".pptx" in ai_assistant._ATTACH_KINDS
    info = ai_assistant._uploaded_file_info(7, _presentacion())
    assert info is not None, "el servidor rechazaba el archivo"
    assert info["kind"] == "pptx"
    assert info["filename"] == "Resultados OQC.pptx"


def test_el_texto_tablas_y_notas_de_la_presentacion_llegan_al_modelo(mundo):
    ref = _presentacion()
    partes = ai_assistant._attachment_input_parts(
        7, ref, ai_assistant._uploaded_file_info(7, ref))
    texto = partes[0]["text"]

    assert "--- diapositiva 1 ---" in texto
    assert "Resultados OQC Agosto 2026" in texto
    # La tabla se vuelca fila por fila.
    assert "Linea | Defectos" in texto
    assert "M1 | 37" in texto
    # Y las notas del presentador, donde suele estar el porque.
    assert "[notas del presentador] M1 subio por el cambio de molde." in texto


def test_pregunta_normal_sobre_una_presentacion(client, mundo):
    ref = _presentacion()
    mundo.guion = [["M1 tuvo 37 defectos y M4 solo 12; las notas lo atribuyen al cambio de molde."]]
    cuerpo = _turno(client, "que dice esta presentacion?", file_refs=[ref])

    assert "37 defectos" in cuerpo
    assert mundo.ejecuciones == [], "una presentacion ajena al MES no usa reportes"
    assert mundo.escrituras == []


def test_el_cliente_ofrece_pptx_igual_que_el_servidor():
    """El catalogo del JS y el accept del input deben seguir al backend."""
    import pathlib as _pl

    js = _pl.Path("app/static/js/ai-assistant.js").read_text(encoding="utf-8")
    html = _pl.Path("app/templates/components/ai_assistant.html").read_text(encoding="utf-8")
    assert "'.pptx'" in js
    assert ".pptx" in html
    for extension in ai_assistant._ATTACH_KINDS:
        assert "'" + extension + "'" in js, extension


# ---------------------------------------------------------------------------
# Generar PowerPoint desde datos de adjuntos.
# ---------------------------------------------------------------------------
# Se arman con la MISMA normalizacion que usa la herramienta: el modelo manda
# los numeros como texto y _tabla_a_filas los convierte para que haya grafica.
COLUMNAS_PPT, FILAS_PPT = ai_assistant._tabla_a_filas(
    ["Persona", "Faltas"],
    [["PERSONA 001", "16"], ["PERSONA 002", "14"], ["PERSONA 003", "13"]],
)


def test_se_puede_construir_una_presentacion_desde_una_tabla_suelta():
    import io as _io

    from pptx import Presentation

    from app.api.portal.ai_artifacts import build_table_powerpoint

    data = build_table_powerpoint(
        title="Faltas 2026", columns=COLUMNAS_PPT, rows=FILAS_PPT)
    pres = Presentation(_io.BytesIO(data))

    assert len(pres.slides) >= 4
    todo = " ".join(
        sh.text_frame.text for d in pres.slides for sh in d.shapes
        if getattr(sh, "has_text_frame", False)
    )
    assert "Faltas 2026" in todo
    # La procedencia no debe decir que salio del MES: no salio.
    assert "no provienen del MES" in todo
    assert "None" not in todo, "quedaron campos sin llenar en la portada"
    assert any(sh.has_chart for d in pres.slides for sh in d.shapes), "falta la grafica"


def test_la_tool_de_presentacion_se_ofrece_junto_a_la_de_excel():
    # El schema existe y describe el caso de uso.
    schema = ai_assistant._table_pptx_tool_schema()
    assert schema["name"] == "powerpoint_desde_tabla"
    assert "presentación" in schema["description"]
    assert "Nunca digas que no tienes herramienta" in schema["description"]


def test_pedir_una_presentacion_genera_el_archivo(client, mundo, monkeypatch):
    """Antes: "No tengo una herramienta habilitada para convertir... en PowerPoint"."""
    registrados = []

    def _fake_register(**kwargs):
        registrados.append(kwargs)
        return {"id": "art-ppt", "type": "pptx", "title": kwargs.get("title"),
                "filename": kwargs.get("filename"), "row_count": kwargs.get("row_count")}

    monkeypatch.setattr(ai_assistant, "register_file_artifact", _fake_register)
    monkeypatch.setattr(ai_assistant, "check_quota", lambda *_a, **_k: (True, None, {}, {}))

    ref = _matriz_asistencia(personas=40, dias=90)
    mundo.guion = [[
        {"tool": "powerpoint_desde_tabla",
         "args": {"titulo": "Faltas y retardos 2026",
                  "columnas": ["Persona", "Faltas"],
                  "filas": [["PERSONA 001", "16"], ["PERSONA 002", "14"]]}},
        "Listo, adjunte la presentacion.",
    ]]
    _turno(client, "genera una presentacion de eso", file_refs=[ref])

    assert registrados, "no se genero la presentacion"
    assert registrados[0]["filename"].endswith(".pptx")
    assert registrados[0]["row_count"] == 2
    assert len(registrados[0]["data"]) > 20000, "el pptx debe traer contenido"


def _texto_de(data: bytes) -> str:
    import io as _io

    from pptx import Presentation

    pres = Presentation(_io.BytesIO(data))
    return " ".join(
        sh.text_frame.text for d in pres.slides for sh in d.shapes
        if getattr(sh, "has_text_frame", False)
    )


def test_la_presentacion_de_adjuntos_no_dice_que_los_datos_son_del_mes():
    """La plantilla afirmaba "datos consultados del MES" en TODA presentacion."""
    from app.api.portal.ai_artifacts import build_table_powerpoint

    texto = _texto_de(build_table_powerpoint(
        title="Faltas 2026", columns=COLUMNAS_PPT, rows=FILAS_PPT))

    assert "no provienen del MES" in texto
    assert "datos consultados del MES" not in texto
    assert "Fuente MES" not in texto, "la etiqueta tambien afirmaba procedencia"
    assert "Archivos adjuntos al chat" in texto


def test_la_presentacion_de_un_reporte_mes_si_lo_dice(tmp_path):
    """El comportamiento de siempre no cambia: sin bandera, sigue siendo MES."""
    from app.api.portal.ai_artifacts import _build_powerpoint

    destino = tmp_path / "reporte.pptx"
    _build_powerpoint(
        {"rows": FILAS_PPT, "columns": COLUMNAS_PPT, "filters": {},
         "source": "Control de material", "report": "warehouse_analysis", "title": "Reporte"},
        "Reporte", "es", destino,
    )
    texto = _texto_de(destino.read_bytes())

    assert "Resultados generados exclusivamente con datos consultados del MES." in texto
    assert "Fuente MES: Control de material" in texto


# ---------------------------------------------------------------------------
# Imagenes DENTRO de contenedores: el modelo debe poder verlas.
# ---------------------------------------------------------------------------
def _png(color=(200, 30, 30), tam=(48, 32)) -> bytes:
    import io as _io

    from PIL import Image

    buffer = _io.BytesIO()
    Image.new("RGB", tam, color).save(buffer, format="PNG")
    return buffer.getvalue()


def _zip_con_imagenes(ref="zip001", conversacion=7):
    import zipfile as _zip

    carpeta = ai_assistant._upload_root() / str(conversacion)
    carpeta.mkdir(parents=True, exist_ok=True)
    destino = carpeta / (ref + ".zip")
    with _zip.ZipFile(destino, "w") as z:
        z.writestr("notas.txt", "Defecto encontrado en linea M1.")
        z.writestr("foto_defecto.png", _png())
        z.writestr("foto_pieza.jpg", _png(color=(20, 90, 200)))
    (carpeta / (ref + ".name")).write_text("evidencias.zip", "utf-8")
    return ref


def test_las_imagenes_dentro_de_un_zip_llegan_al_modelo(mundo):
    """Antes se omitian: un zip de evidencias fotograficas llegaba vacio."""
    ref = _zip_con_imagenes()
    partes = ai_assistant._attachment_input_parts(
        7, ref, ai_assistant._uploaded_file_info(7, ref))

    imagenes = [p for p in partes if p["type"] == "input_image"]
    textos = " ".join(p["text"] for p in partes if p["type"] == "input_text")

    assert len(imagenes) == 2, "las dos fotos deben viajar"
    assert all(p["image_url"].startswith("data:image/") for p in imagenes)
    # El texto del zip sigue llegando, y se dice que fotos son.
    assert "Defecto encontrado en linea M1." in textos
    assert "foto_defecto.png" in textos and "foto_pieza.jpg" in textos


def test_un_zip_solo_de_fotos_ya_no_se_reporta_como_ilegible(mundo):
    import zipfile as _zip

    carpeta = ai_assistant._upload_root() / "7"
    carpeta.mkdir(parents=True, exist_ok=True)
    with _zip.ZipFile(carpeta / "solofotos.zip", "w") as z:
        z.writestr("a.png", _png())
    (carpeta / "solofotos.name").write_text("solo fotos.zip", "utf-8")

    partes = ai_assistant._attachment_input_parts(
        7, "solofotos", ai_assistant._uploaded_file_info(7, "solofotos"))
    textos = " ".join(p["text"] for p in partes if p["type"] == "input_text")

    assert any(p["type"] == "input_image" for p in partes)
    assert "no trae archivos" not in textos


def test_las_imagenes_de_una_presentacion_llegan_al_modelo(mundo):
    """Una diapositiva que es una captura llegaba sin nada que mirar."""
    import io as _io

    from pptx import Presentation
    from pptx.util import Inches

    carpeta = ai_assistant._upload_root() / "7"
    carpeta.mkdir(parents=True, exist_ok=True)
    pres = Presentation()
    hoja = pres.slides.add_slide(pres.slide_layouts[5])
    hoja.shapes.title.text = "Evidencia del defecto"
    hoja.shapes.add_picture(_io.BytesIO(_png(tam=(120, 80))), Inches(1), Inches(2))
    pres.save(carpeta / "conimg.pptx")
    (carpeta / "conimg.name").write_text("Evidencia.pptx", "utf-8")

    partes = ai_assistant._attachment_input_parts(
        7, "conimg", ai_assistant._uploaded_file_info(7, "conimg"))

    assert any(p["type"] == "input_image" for p in partes), "la captura debe viajar"
    textos = " ".join(p["text"] for p in partes if p["type"] == "input_text")
    assert "Evidencia del defecto" in textos
    assert "imagenes suyas ademas del texto" in textos


def test_una_presentacion_sin_imagenes_sigue_viajando_como_texto(mundo):
    ref = _presentacion()
    partes = ai_assistant._attachment_input_parts(
        7, ref, ai_assistant._uploaded_file_info(7, ref))
    assert [p["type"] for p in partes] == ["input_text"]
    assert "M1 | 37" in partes[0]["text"]


def test_una_imagen_enorme_dentro_del_zip_se_omite_sin_romper(mundo):
    """Tope por imagen: una foto gigante no debe tumbar el turno completo."""
    import zipfile as _zip

    carpeta = ai_assistant._upload_root() / "7"
    carpeta.mkdir(parents=True, exist_ok=True)
    with _zip.ZipFile(carpeta / "pesado.zip", "w") as z:
        z.writestr("chica.png", _png())
        z.writestr("gigante.png", b"\x89PNG\r\n\x1a\n" + b"0" * (5 * 1024 * 1024))
    (carpeta / "pesado.name").write_text("pesado.zip", "utf-8")

    partes = ai_assistant._attachment_input_parts(
        7, "pesado", ai_assistant._uploaded_file_info(7, "pesado"))
    assert len([p for p in partes if p["type"] == "input_image"]) == 1


def test_el_chat_acepta_pegar_con_ctrl_v():
    """Ctrl+V debe entrar por la misma ruta que el drag & drop."""
    import pathlib as _pl

    js = _pl.Path("app/static/js/ai-assistant.js").read_text(encoding="utf-8")

    # El listener existe y reutiliza uploadPlanFiles.
    assert "addEventListener('paste'" in js
    assert "archivosDelPortapapeles" in js
    indice = js.index("addEventListener('paste'")
    bloque = js[indice:indice + 400]
    assert "uploadPlanFiles" in bloque
    # Si no hay archivos, el pegado de texto normal no se toca.
    assert "if (!archivos.length) return;" in bloque
    assert bloque.index("if (!archivos.length) return;") < bloque.index("preventDefault")
    # Una captura sin nombre util se renombra con fecha.
    assert "nombrarPegado" in js
    assert "captura-" in js
    # Los dos caminos del portapapeles segun navegador.
    assert "data?.files" in js and "item.getAsFile()" in js
    # Y se anuncia en el placeholder de los tres idiomas.
    import re as _re

    placeholders = _re.findall(r"placeholder:'([^']*)'", js)
    assert len(placeholders) == 3, placeholders
    assert all("Ctrl+V" in texto for texto in placeholders), placeholders


# ---------------------------------------------------------------------------
# Estilo: corporativo por defecto, propio solo si el usuario lo pide.
# ---------------------------------------------------------------------------
def _abrir_pptx(data: bytes):
    import io as _io

    from pptx import Presentation

    return Presentation(_io.BytesIO(data))


def _abrir_xlsx(data: bytes):
    import io as _io

    from openpyxl import load_workbook

    return load_workbook(_io.BytesIO(data)).active


def test_por_defecto_sale_el_estilo_corporativo(mundo):
    from app.api.portal.ai_artifacts import build_table_excel, build_table_powerpoint, plantilla_pptx

    pres = _abrir_pptx(build_table_powerpoint(
        title="Inventario", columns=COLUMNAS_PPT, rows=FILAS_PPT))
    base = _abrir_pptx(plantilla_pptx().read_bytes())

    # Hereda el tamano de la plantilla ISEMM, no un 16:9 inventado.
    assert (pres.slide_width, pres.slide_height) == (base.slide_width, base.slide_height)
    titulo = [s for s in list(pres.slides)[0].shapes if getattr(s, "has_text_frame", False)][0]
    fuente = titulo.text_frame.paragraphs[0].font
    assert fuente.name == "LG Smart UI Regular"
    assert str(fuente.color.rgb) == "1F497D"

    hoja = _abrir_xlsx(build_table_excel(
        title="Inventario", columns=COLUMNAS_PPT, rows=FILAS_PPT))
    assert hoja["A1"].font.name == "LG Smart UI Regular"
    assert str(hoja["A1"].fill.start_color.rgb).endswith("000000"), "encabezado negro"
    # Sin el azul de Excel: la tabla va en blanco y negro.
    estilo_tabla = list(hoja.tables.values())[0].tableStyleInfo.name
    assert "Medium2" not in estilo_tabla, estilo_tabla
    assert estilo_tabla == "TableStyleLight1"
    # La tipografia tambien en los datos, no solo en el encabezado.
    assert hoja["A2"].font.name == "LG Smart UI Regular"


def test_el_texto_del_encabezado_se_adapta_al_relleno(mundo):
    """Con el gris LG el blanco de siempre habria quedado ilegible."""
    from app.api.portal.ai_artifacts import build_table_excel

    claro = _abrir_xlsx(build_table_excel(
        title="x", columns=COLUMNAS_PPT, rows=FILAS_PPT,
        estilo={"relleno_encabezado": "D9D9D9"}))
    oscuro = _abrir_xlsx(build_table_excel(
        title="x", columns=COLUMNAS_PPT, rows=FILAS_PPT))

    assert str(claro["A1"].font.color.rgb).endswith("000000"), "gris claro pide texto negro"
    assert str(oscuro["A1"].font.color.rgb).endswith("FFFFFF"), "negro pide texto blanco"


def test_el_usuario_puede_pedir_otro_estilo(mundo):
    from app.api.portal.ai_artifacts import build_table_excel, build_table_powerpoint

    estilo = {"fuente": "Arial", "color_titulo": "1B7F3B", "relleno_encabezado": "1B7F3B"}
    pres = _abrir_pptx(build_table_powerpoint(
        title="Inventario", columns=COLUMNAS_PPT, rows=FILAS_PPT, estilo=estilo))
    titulo = [s for s in list(pres.slides)[0].shapes if getattr(s, "has_text_frame", False)][0]
    fuente = titulo.text_frame.paragraphs[0].font

    assert fuente.name == "Arial"
    assert str(fuente.color.rgb) == "1B7F3B"
    hoja = _abrir_xlsx(build_table_excel(
        title="x", columns=COLUMNAS_PPT, rows=FILAS_PPT, estilo=estilo))
    assert hoja["A1"].font.name == "Arial"


def test_un_estilo_invalido_no_rompe_el_archivo(mundo):
    """El modelo no puede tumbar la generacion mandando basura."""
    from app.api.portal.ai_artifacts import resolver_estilo

    activo = resolver_estilo({
        "color_titulo": "no-es-un-color",
        "relleno_encabezado": "#ZZZZZZ",
        "fuente": "",
        "clave_inventada": "lo que sea",
    })
    assert activo["color_titulo"] == "1F497D", "un hex invalido cae al corporativo"
    assert activo["relleno_encabezado"] == "000000"
    assert activo["fuente"] == "LG Smart UI Regular"
    assert "clave_inventada" not in activo


def test_el_prompt_le_dice_que_el_estilo_corporativo_es_el_default():
    prompt = ai_openai.build_instructions({"language": "es"})
    assert "estilo corporativo" in prompt
    assert "deja el parámetro estilo en null" in prompt
    assert "Sólo mandalo cuando el usuario pida explícitamente" in prompt


# ---------------------------------------------------------------------------
# Los schemas deben ser validos para OpenAI: uno malo tira TODA peticion con 400.
# ---------------------------------------------------------------------------
def _todos_los_schemas():
    """Cada schema de herramienta que el asistente puede declarar."""
    hojas = ["Activos", "LUNES"]
    return [
        ai_assistant._table_excel_tool_schema(),
        ai_assistant._table_pptx_tool_schema(),
        ai_assistant._excel_edit_tool_schema(),
        ai_assistant._excel_leer_tool_schema(hojas),
        ai_assistant._excel_contar_tool_schema(hojas),
        ai_assistant._pdf_buscar_tool_schema("manual.pdf"),
        ai_assistant._pdf_leer_tool_schema("manual.pdf"),
        ai_assistant._imagen_tool_schema(),
    ]


def _revisar_objeto(nodo, ruta, fallas):
    """En strict, required debe listar EXACTAMENTE las claves de properties."""
    if not isinstance(nodo, dict):
        return
    if nodo.get("type") == "object" or "properties" in nodo:
        props = nodo.get("properties")
        if props is None:
            fallas.append(f"{ruta}: objeto sin properties")
            return
        requeridas = set(nodo.get("required") or [])
        sobran = requeridas - set(props)
        faltan = set(props) - requeridas
        if sobran:
            fallas.append(f"{ruta}: required nombra claves que no existen: {sorted(sobran)}")
        if faltan:
            fallas.append(f"{ruta}: properties sin declarar en required: {sorted(faltan)}")
        if nodo.get("additionalProperties") is not False:
            fallas.append(f"{ruta}: falta additionalProperties=False")
        for clave, valor in props.items():
            _revisar_objeto(valor, f"{ruta}.{clave}", fallas)
    if nodo.get("type") == "array":
        _revisar_objeto(nodo.get("items") or {}, f"{ruta}[]", fallas)


def test_los_schemas_de_herramientas_son_validos_para_openai():
    """Un 'estilo' fuera de properties tiraba todas las peticiones con 400.

    El error no lo cachaban los tests de los constructores: hay que revisar el
    schema que se le manda al modelo, no solo el archivo que sale al final.
    """
    fallas = []
    for schema in _todos_los_schemas():
        assert schema.get("type") == "function", schema.get("name")
        assert schema.get("name"), "toda herramienta necesita nombre"
        if schema.get("strict"):
            _revisar_objeto(schema["parameters"], schema["name"], fallas)
    assert not fallas, "\n".join(fallas)


def test_los_schemas_viajan_como_json():
    """Si algo no es serializable el fallo aparece hasta la llamada real."""
    for schema in _todos_los_schemas():
        json.dumps(schema, ensure_ascii=False)


def _lineas_de_la_plantilla():
    """Las lineas horizontales que dibuja el layout corporativo."""
    from pptx import Presentation

    from app.api.portal.ai_artifacts import plantilla_pptx

    base = Presentation(str(plantilla_pptx()))
    return sorted(
        sh.top for sh in base.slide_layouts[0].shapes
        if sh.height == 0 and str(sh.shape_type).startswith("LINE")
    )


def test_nada_de_lo_generado_pisa_la_decoracion_de_la_plantilla(mundo):
    """El titulo salia encima de la linea verde del encabezado ISEMM."""
    from app.api.portal.ai_artifacts import build_table_powerpoint

    lineas = _lineas_de_la_plantilla()
    assert lineas, "la plantilla deberia traer sus lineas"

    pres = _abrir_pptx(build_table_powerpoint(
        title="Faltas 2026", columns=COLUMNAS_PPT, rows=FILAS_PPT))

    encimados = [
        (i, sh.top, sh.height, linea)
        for i, d in enumerate(pres.slides, 1)
        for sh in d.shapes
        for linea in lineas
        if sh.top < linea < sh.top + sh.height
    ]
    assert not encimados, encimados


def test_no_se_agrega_un_segundo_logo_sobre_el_de_la_plantilla(mundo):
    """La plantilla ya trae su marca; el logo ILSAN extra la duplicaba."""
    from app.api.portal.ai_artifacts import build_table_powerpoint

    pres = _abrir_pptx(build_table_powerpoint(
        title="Faltas 2026", columns=COLUMNAS_PPT, rows=FILAS_PPT))
    imagenes = [sh for d in pres.slides for sh in d.shapes if sh.shape_type == 13]
    assert not imagenes, "con plantilla, las imagenes las pone la plantilla"


def test_sin_plantilla_se_conserva_el_formato_anterior(tmp_path, monkeypatch):
    """Quitar la plantilla no debe dejar el generador sin diseno."""
    from app.api.portal import ai_artifacts

    monkeypatch.setenv("AI_PPTX_TEMPLATE", str(tmp_path / "no_existe.pptx"))
    destino = tmp_path / "r.pptx"
    ai_artifacts._build_powerpoint(
        {"rows": FILAS_PPT, "columns": COLUMNAS_PPT, "filters": {},
         "source": "MES", "report": "r", "title": "Reporte"},
        "Reporte", "es", destino,
    )
    pres = _abrir_pptx(destino.read_bytes())
    assert round(pres.slide_width / pres.slide_height, 2) == round(16 / 9, 2)
    assert any(sh.shape_type == 13 for d in pres.slides for sh in d.shapes), "vuelve el logo"


def test_la_barra_de_cuota_muestra_el_acumulado_del_dia_no_la_respuesta(client, mundo, monkeypatch):
    """Decia "22,532 / 20,000,000" con 1,121,778 tokens gastados en el dia."""
    monkeypatch.setattr(ai_assistant, "check_quota",
                        lambda *_a, **_k: (True, None,
                                           {"input_tokens": 1_100_000, "output_tokens": 21_778,
                                            "request_count": 31, "artifact_count": 4},
                                           {"daily_token_limit": 20_000_000,
                                            "daily_request_limit": 500,
                                            "daily_artifact_limit": 100}))
    mundo.guion = [["Listo."]]
    cuerpo = _turno(client, "hola")

    assert "usage_diaria" in cuerpo, "debe emitirse el acumulado del dia"
    evento = [l for l in cuerpo.splitlines() if l.startswith("data:") and "1100000" in l]
    assert evento, cuerpo[-500:]
    datos = json.loads(evento[0][len("data:"):])
    assert datos["input_tokens"] == 1_100_000
    assert datos["limits"]["daily_token_limit"] == 20_000_000


def test_el_cliente_usa_el_evento_diario_para_la_cuota():
    import pathlib as _pl

    js = _pl.Path("app/static/js/ai-assistant.js").read_text(encoding="utf-8")
    assert "if (event === 'usage_diaria')" in js
    # El evento por respuesta ya no debe pisar la barra.
    assert "if (event === 'usage') this.updateUsage" not in js


# ---------------------------------------------------------------------------
# Generacion de imagenes con IA.
# ---------------------------------------------------------------------------
def _png_falso() -> bytes:
    return _png(color=(10, 120, 200), tam=(64, 64))


def test_el_modelo_de_imagen_se_valida_contra_la_lista(monkeypatch):
    """Un nombre inventado o retirado no debe tumbar la generacion."""
    from app.api.portal import ai_openai

    llamadas = {}

    class _Imagenes:
        def generate(self, **kw):
            llamadas.update(kw)
            import base64
            from types import SimpleNamespace
            return SimpleNamespace(
                data=[SimpleNamespace(b64_json=base64.b64encode(_png_falso()).decode())],
                usage=None)

    monkeypatch.setattr(ai_openai, "_client",
                        lambda: type("C", (), {"images": _Imagenes()})())
    monkeypatch.setenv("AI_SAFETY_HMAC_KEY", "x")

    ai_openai.generate_image(prompt="un icono", model="modelo-que-no-existe")
    assert llamadas["model"] == ai_openai.image_model(), "cae al de por defecto"

    ai_openai.generate_image(prompt="un icono", model="gpt-image-2.5-flare")
    assert llamadas["model"] == "gpt-image-2.5-flare", "un modelo valido si se respeta"


def test_la_herramienta_ofrece_los_modelos_de_la_cuenta():
    schema = ai_assistant._imagen_tool_schema()
    opciones = schema["parameters"]["properties"]["modelo"]["enum"]

    assert "gpt-image-2" in opciones
    assert "gpt-image-2.5-flare" in opciones and "gpt-image-2.5-sunburst" in opciones
    assert None in opciones, "debe poder no elegir"
    # Y se le pide que justifique la eleccion.
    assert "explica" in schema["parameters"]["properties"]["modelo"]["description"]
    assert "cuesta dinero" in schema["description"]


def test_generar_imagen_deja_un_artefacto_descargable(client, mundo, monkeypatch):
    from app.api.portal import ai_openai

    registrados = []
    monkeypatch.setattr(ai_assistant, "register_file_artifact",
                        lambda **kw: registrados.append(kw) or {
                            "id": "img-1", "type": "png", "title": kw.get("title"),
                            "filename": kw.get("filename")})
    monkeypatch.setattr(ai_assistant, "check_quota", lambda *_a, **_k: (True, None, {}, {}))
    monkeypatch.setattr(ai_assistant, "generate_image",
                        lambda **kw: (_png_falso(), {"model": kw.get("model") or "gpt-image-2",
                                                     "input_tokens": 0, "output_tokens": 0}))

    mundo.guion = [[
        {"tool": "generar_imagen",
         "args": {"descripcion": "icono de una laptop", "tamano": "1024x1024",
                  "calidad": "low", "modelo": "gpt-image-2.5-flare"}},
        "Listo, la imagen quedo adjunta.",
    ]]
    _turno(client, "generame un icono de una laptop")

    assert registrados, "no se registro la imagen"
    assert registrados[0]["filename"].endswith(".png")
    assert registrados[0]["data"][:8] == b"\x89PNG\r\n\x1a\n"
    entregado = json.dumps(mundo.vistas_por_el_modelo[-1], ensure_ascii=False, default=str)
    assert "gpt-image-2.5-flare" in entregado, "se le informa que modelo se uso"
    assert "img-1" in entregado, "y el id para poder incrustarla"


def test_la_imagen_solo_se_incrusta_si_se_pasa_su_id(mundo):
    from app.api.portal.ai_artifacts import build_table_powerpoint

    sin = _abrir_pptx(build_table_powerpoint(
        title="x", columns=COLUMNAS_PPT, rows=FILAS_PPT))
    con = _abrir_pptx(build_table_powerpoint(
        title="x", columns=COLUMNAS_PPT, rows=FILAS_PPT, imagen=_png_falso()))

    assert not any(sh.shape_type == 13 for d in sin.slides for sh in d.shapes)
    assert any(sh.shape_type == 13 for d in con.slides for sh in d.shapes)
    assert len(con.slides) == len(sin.slides) + 1, "la imagen trae su propia diapositiva"


def test_la_imagen_en_excel_va_en_su_propia_hoja(mundo):
    from app.api.portal.ai_artifacts import build_table_excel

    import io as _io

    from openpyxl import load_workbook

    libro = load_workbook(_io.BytesIO(build_table_excel(
        title="x", columns=COLUMNAS_PPT, rows=FILAS_PPT, imagen=_png_falso())))
    assert "Imagen" in libro.sheetnames, "no debe tapar los datos"
    assert libro["Imagen"]._images, "la imagen debe estar incrustada"


def test_la_imagen_generada_aparece_como_tarjeta_en_el_chat(client, mundo, monkeypatch):
    """El archivo se creaba pero el panel no lo mostraba: faltaba artifact_ready."""
    monkeypatch.setattr(ai_assistant, "register_file_artifact",
                        lambda **kw: {"id": "img-9", "type": "png", "status": "ready",
                                      "title": kw.get("title"), "filename": kw.get("filename"),
                                      "download_url": "/api/ai/artifacts/img-9/download"})
    monkeypatch.setattr(ai_assistant, "check_quota", lambda *_a, **_k: (True, None, {}, {}))
    monkeypatch.setattr(ai_assistant, "generate_image",
                        lambda **kw: (_png_falso(), {"model": "gpt-image-2",
                                                     "input_tokens": 0, "output_tokens": 0}))

    mundo.guion = [[
        {"tool": "generar_imagen",
         "args": {"descripcion": "cartel navideno", "tamano": "1024x1536",
                  "calidad": "medium", "modelo": None}},
        "Lista, quedo adjunta.",
    ]]
    cuerpo = _turno(client, "genera una imagen de navidad")

    assert "artifact_ready" in cuerpo, "el panel necesita este evento para pintar la tarjeta"
    assert "img-9" in cuerpo
    assert "/api/ai/artifacts/img-9/download" in cuerpo, "y su enlace de descarga"


# ---------------------------------------------------------------------------
# Autocompacto del historial.
# ---------------------------------------------------------------------------
def test_una_charla_corta_no_se_recorta():
    mensajes = [{"role": "user", "content": f"pregunta {i}"} for i in range(20)]
    ventana, cortados = ai_assistant._compactar_historial(mensajes)
    assert cortados == 0 and len(ventana) == 20


def test_los_mensajes_pesados_se_recortan_por_los_mas_viejos():
    """Un volcado de adjunto en el historial llenaba el contexto solo."""
    pesado = "x" * 20000
    mensajes = [{"role": "user", "content": pesado} for _ in range(10)]
    mensajes.append({"role": "user", "content": "la pregunta de ahora"})

    ventana, cortados = ai_assistant._compactar_historial(mensajes)

    assert cortados > 0, "algo se debe cortar"
    peso = sum(ai_assistant._peso_mensaje(m) for m in ventana)
    assert peso <= ai_assistant._HISTORIAL_MAX_CHARS or len(ventana) == ai_assistant._HISTORIAL_MIN_MENSAJES
    # Se corta por delante: lo ultimo siempre sobrevive.
    assert ventana[-1]["content"] == "la pregunta de ahora"


def test_nunca_se_queda_sin_los_ultimos_mensajes():
    """Aunque cada mensaje exceda el presupuesto, el hilo inmediato se conserva."""
    enorme = "y" * 200000
    mensajes = [{"role": "user", "content": enorme} for _ in range(12)]
    ventana, _cortados = ai_assistant._compactar_historial(mensajes)
    assert len(ventana) == ai_assistant._HISTORIAL_MIN_MENSAJES


def test_el_peso_cuenta_los_adjuntos_no_solo_el_texto():
    """El contenido puede ser una lista de partes; contarla como 0 era el error."""
    con_partes = {"role": "user", "content": [
        {"type": "input_text", "text": "z" * 5000},
        {"type": "input_text", "text": "z" * 5000},
    ]}
    assert ai_assistant._peso_mensaje(con_partes) > 10000


def test_el_resumen_cubre_exactamente_lo_que_no_se_envia():
    """Antes el resumen guardaba desde el mensaje 12 y se enviaban 20: los
    mensajes 13 a 20 viajaban duplicados, verbatim y resumidos."""
    import pathlib as _pl

    codigo = _pl.Path("app/api/portal/ai_assistant.py").read_text(encoding="utf-8")
    assert "keep_recent=12" not in codigo
    assert codigo.count("keep_recent=_HISTORIAL_MAX_MENSAJES") == 3


# ---------------------------------------------------------------------------
# Vista previa de imagenes en la tarjeta.
# ---------------------------------------------------------------------------
def _js():
    import pathlib as _pl

    return _pl.Path("app/static/js/ai-assistant.js").read_text(encoding="utf-8")


def test_la_tarjeta_de_imagen_muestra_vista_previa():
    js = _js()
    assert "ai-artifact-preview" in js
    assert "'png','jpg','jpeg','webp','gif'" in js
    # Se pide en linea para no ensuciar la auditoria de descargas.
    assert "'?inline=1'" in js
    # Y "0 filas" no aplica a una imagen.
    indice = js.index("const esImagen")
    assert "esImagen" in js[indice:indice + 600]
    assert "rows" in js[indice:indice + 600]


def test_el_css_de_la_vista_previa_existe():
    import pathlib as _pl

    css = _pl.Path("app/static/css/ai-assistant.css").read_text(encoding="utf-8")
    assert ".ai-artifact-preview{" in css
    assert "object-fit:contain" in css, "no debe deformar la imagen"
    assert "max-height" in css, "una imagen alta no debe empujar el chat"


def test_ver_en_linea_no_cuenta_como_descarga():
    """Pintar la miniatura registraba una descarga falsa en la auditoria."""
    import pathlib as _pl

    codigo = _pl.Path("app/api/portal/ai_assistant.py").read_text(encoding="utf-8")
    indice = codigo.index("def download_artifact")
    bloque = codigo[indice:indice + 2600]

    assert 'request.args.get("inline")' in bloque
    assert 'startswith("image/")' in bloque, "solo imagenes en linea"
    assert "as_attachment=not en_linea" in bloque
    # El registro de descarga queda condicionado.
    assert bloque.index("if not en_linea:") < bloque.index("DESCARGAR_ARTEFACTO")


# ---------------------------------------------------------------------------
# El logo real como referencia, no uno inventado.
# ---------------------------------------------------------------------------
class _ImagenesEspia:
    """Registra si se llamo generate o edit y con que."""

    def __init__(self, rechaza_fidelity=False):
        self.rechaza_fidelity = rechaza_fidelity
        self.llamadas = []

    def _respuesta(self):
        import base64
        from types import SimpleNamespace

        return SimpleNamespace(
            data=[SimpleNamespace(b64_json=base64.b64encode(_png_falso()).decode())],
            usage=None)

    def generate(self, **kw):
        self.llamadas.append(("generate", kw))
        return self._respuesta()

    def edit(self, **kw):
        if self.rechaza_fidelity and "input_fidelity" in kw:
            raise RuntimeError("The model does not support the 'input_fidelity' parameter.")
        self.llamadas.append(("edit", kw))
        return self._respuesta()


def _espiar(monkeypatch, espia):
    from app.api.portal import ai_openai

    monkeypatch.setattr(ai_openai, "_client",
                        lambda: type("C", (), {"images": espia})())
    monkeypatch.setenv("AI_SAFETY_HMAC_KEY", "x")
    return ai_openai


def test_el_logo_corporativo_es_el_archivo_real():
    from app.api.portal.ai_openai import logo_corporativo

    referencia = logo_corporativo()
    assert referencia is not None, "no se encontro app/static/images/ilsan-logo.png"
    nombre, datos = referencia
    assert nombre == "ilsan-logo.png"
    assert datos[:8] == b"\x89PNG\r\n\x1a\n"
    assert len(datos) > 1000


def test_sin_referencias_se_genera_desde_cero(monkeypatch):
    espia = _ImagenesEspia()
    ai_openai = _espiar(monkeypatch, espia)
    ai_openai.generate_image(prompt="un cartel")
    assert espia.llamadas[0][0] == "generate"


def test_con_logo_se_manda_el_archivo_como_referencia(monkeypatch):
    """Describir el logo con palabras producia una version falsa parecida."""
    espia = _ImagenesEspia()
    ai_openai = _espiar(monkeypatch, espia)
    ai_openai.generate_image(prompt="cartel navideno",
                             referencias=[ai_openai.logo_corporativo()])

    modo, kw = espia.llamadas[0]
    assert modo == "edit", "con referencias se edita, no se genera de cero"
    assert len(kw["image"]) == 1
    nombre, datos, mime = kw["image"][0]
    assert nombre == "ilsan-logo.png" and mime == "image/png"
    assert datos[:8] == b"\x89PNG\r\n\x1a\n", "viajan los bytes reales del logo"


def test_si_el_modelo_rechaza_input_fidelity_se_edita_igual(monkeypatch):
    """gpt-image-2 no acepta ese parametro; no debe tumbar la edicion."""
    espia = _ImagenesEspia(rechaza_fidelity=True)
    ai_openai = _espiar(monkeypatch, espia)

    ai_openai.generate_image(prompt="x", referencias=[("logo.png", _png_falso())])

    assert len(espia.llamadas) == 1, "solo la llamada buena queda registrada"
    modo, kw = espia.llamadas[0]
    assert modo == "edit" and "input_fidelity" not in kw


def test_un_error_distinto_no_se_traga(monkeypatch):
    """El reintento es solo para input_fidelity, no para cualquier fallo."""
    class _Rota(_ImagenesEspia):
        def edit(self, **kw):
            raise RuntimeError("billing hard limit reached")

    ai_openai = _espiar(monkeypatch, _Rota())
    try:
        ai_openai.generate_image(prompt="x", referencias=[("logo.png", _png_falso())])
        raise AssertionError("deberia propagar el error")
    except RuntimeError as exc:
        assert "billing" in str(exc)


def test_editar_una_imagen_previa_la_manda_de_base(client, mundo, monkeypatch):
    from app.api.portal import ai_openai

    recibidas = {}
    monkeypatch.setattr(ai_assistant, "register_file_artifact",
                        lambda **kw: {"id": "img-2", "type": "png", "status": "ready",
                                      "title": kw.get("title"), "filename": kw.get("filename"),
                                      "download_url": "/d"})
    monkeypatch.setattr(ai_assistant, "check_quota", lambda *_a, **_k: (True, None, {}, {}))
    monkeypatch.setattr(ai_assistant, "_bytes_de_artefacto",
                        lambda pid: b"\x89PNG\r\n\x1a\nANTERIOR" if pid else None)
    monkeypatch.setattr(ai_assistant, "generate_image",
                        lambda **kw: recibidas.update(kw) or (_png_falso(), {"model": "gpt-image-2",
                                                                            "input_tokens": 0,
                                                                            "output_tokens": 0}))
    mundo.guion = [[
        {"tool": "generar_imagen",
         "args": {"descripcion": "anade el logo", "tamano": None, "calidad": None,
                  "modelo": None, "usar_logo": True, "editar_imagen_id": "img-1"}},
        "Listo.",
    ]]
    _turno(client, "anade el logo de ilsan")

    nombres = [n for n, _d in (recibidas.get("referencias") or [])]
    assert nombres == ["anterior.png", "ilsan-logo.png"], nombres


def test_el_prompt_prohibe_dibujar_el_logo_de_memoria():
    prompt = ai_openai.build_instructions({"language": "es"})
    assert "usar_logo=true" in prompt
    assert "Jamás describas el logo con palabras" in prompt
    assert "editar_imagen_id" in prompt
